"""E2E Test Suite for Grounded Vietnamese Analytical Runs: PPTX, XLSX, DOCX.

Verifies:
1. PPTX deck generation with native PowerPoint charts, grounded KPI cards, and exclusion of unverified hallucinated numbers.
2. Multi-sheet XLSX workbook generation with precomputed totals, growth metrics, and formulas.
3. Structured DOCX report generation where every number is strictly traceable to database query results.
"""

import io
import re
from pathlib import Path
import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from dbgpt_analyst.tools import deck, officecli
from dbgpt_analyst.subgraphs.modes.docx_generator import html_to_docx
from dbgpt_analyst.subgraphs.modes.office_writer_subgraph import (
    _grounded_facts,
    _strip_unverified_lines,
    _facts_section,
    _export_excel_file,
)
import markdown as _markdown


# --------------------------------------------------------------------------
# Test Data
# --------------------------------------------------------------------------
PPTX_ROWS = [
    {"quy": "Q1", "kenh": "Online", "doanh_thu": 1_450_000_000, "so_don": 3625, "aov": 400_000},
    {"quy": "Q1", "kenh": "Offline", "doanh_thu": 980_000_000, "so_don": 1960, "aov": 500_000},
    {"quy": "Q2", "kenh": "Online", "doanh_thu": 1_820_000_000, "so_don": 4333, "aov": 420_000},
    {"quy": "Q2", "kenh": "Offline", "doanh_thu": 1_020_000_000, "so_don": 2000, "aov": 510_000},
    {"quy": "Q3", "kenh": "Online", "doanh_thu": 2_150_000_000, "so_don": 4778, "aov": 450_000},
    {"quy": "Q3", "kenh": "Offline", "doanh_thu": 1_050_000_000, "so_don": 2019, "aov": 520_000},
    {"quy": "Q4", "kenh": "Online", "doanh_thu": 3_100_000_000, "so_don": 6458, "aov": 480_000},
    {"quy": "Q4", "kenh": "Offline", "doanh_thu": 1_280_000_000, "so_don": 2370, "aov": 540_000},
]

PPTX_SPEC = {
    "title": "Báo Cáo Phân Tích Doanh Thu Kênh Phân Phối 2025",
    "subtitle": "Đánh giá tăng trưởng các quý và hiệu quả kênh Online so với Offline",
    "source": "Nguồn: Hệ thống ERP & CRM",
    "slides": [
        {
            "kind": "kpi",
            "title": "Chỉ Số Hiệu Suất Tổng Quan",
            "kpis": [
                {"label": "Tổng doanh thu toàn quốc", "metric": "sum:doanh_thu"},
                {"label": "Tổng số lượng đơn hàng", "metric": "sum:so_don"},
                {"label": "Tăng trưởng doanh thu Q4/Q3", "metric": "growth:doanh_thu:quy"},
            ],
        },
        {
            "kind": "chart",
            "title": "Cơ Cấu Doanh Thu Theo Quý & Kênh",
            "chart": "column",
            "category_column": "quy",
            "value_column": "doanh_thu",
            "series_column": "kenh",
        },
        {
            "kind": "table",
            "title": "Bảng Dữ Liệu Bán Hàng Chi Tiết",
            "columns": ["quy", "kenh", "doanh_thu", "so_don", "aov"],
            "limit": 8,
        },
        {
            "kind": "bullets",
            "title": "Nhận Định Chiến Lược & Xu Hướng",
            "bullets": [
                "Kênh Online giữ vai trò động lực tăng trưởng chính qua các quý",
                "Kênh Offline duy trì doanh số ổn định và giá trị đơn hàng trung bình cao",
                "Doanh thu đạt mức kỷ lục 99999999999 đồng vào cuối năm",  # Hallucinated number
            ],
        },
    ],
}

XLSX_ROWS = [
    {"vung_mien": "Miền Bắc", "quy": "Q1", "doanh_thu_thuan": 2_500_000_000, "chi_phi_van_hanh": 1_200_000_000, "loi_nhuan_gop": 1_300_000_000, "so_luong_khach": 12_500},
    {"vung_mien": "Miền Trung", "quy": "Q1", "doanh_thu_thuan": 1_400_000_000, "chi_phi_van_hanh": 750_000_000, "loi_nhuan_gop": 650_000_000, "so_luong_khach": 7_200},
    {"vung_mien": "Miền Nam", "quy": "Q1", "doanh_thu_thuan": 3_200_000_000, "chi_phi_van_hanh": 1_500_000_000, "loi_nhuan_gop": 1_700_000_000, "so_luong_khach": 16_000},
    {"vung_mien": "Miền Bắc", "quy": "Q2", "doanh_thu_thuan": 2_800_000_000, "chi_phi_van_hanh": 1_300_000_000, "loi_nhuan_gop": 1_500_000_000, "so_luong_khach": 13_800},
    {"vung_mien": "Miền Trung", "quy": "Q2", "doanh_thu_thuan": 1_550_000_000, "chi_phi_van_hanh": 800_000_000, "loi_nhuan_gop": 750_000_000, "so_luong_khach": 7_800},
    {"vung_mien": "Miền Nam", "quy": "Q2", "doanh_thu_thuan": 3_600_000_000, "chi_phi_van_hanh": 1_650_000_000, "loi_nhuan_gop": 1_950_000_000, "so_luong_khach": 17_500},
    {"vung_mien": "Miền Bắc", "quy": "Q3", "doanh_thu_thuan": 3_100_000_000, "chi_phi_van_hanh": 1_400_000_000, "loi_nhuan_gop": 1_700_000_000, "so_luong_khach": 15_000},
    {"vung_mien": "Miền Trung", "quy": "Q3", "doanh_thu_thuan": 1_650_000_000, "chi_phi_van_hanh": 820_000_000, "loi_nhuan_gop": 830_000_000, "so_luong_khach": 8_100},
    {"vung_mien": "Miền Nam", "quy": "Q3", "doanh_thu_thuan": 3_950_000_000, "chi_phi_van_hanh": 1_750_000_000, "loi_nhuan_gop": 2_200_000_000, "so_luong_khach": 19_000},
    {"vung_mien": "Miền Bắc", "quy": "Q4", "doanh_thu_thuan": 3_900_000_000, "chi_phi_van_hanh": 1_650_000_000, "loi_nhuan_gop": 2_250_000_000, "so_luong_khach": 18_500},
    {"vung_mien": "Miền Trung", "quy": "Q4", "doanh_thu_thuan": 2_100_000_000, "chi_phi_van_hanh": 950_000_000, "loi_nhuan_gop": 1_150_000_000, "so_luong_khach": 10_200},
    {"vung_mien": "Miền Nam", "quy": "Q4", "doanh_thu_thuan": 4_800_000_000, "chi_phi_van_hanh": 2_000_000_000, "loi_nhuan_gop": 2_800_000_000, "so_luong_khach": 23_000},
]

DOCX_ROWS = [
    {"nganh_hang": "Điện tử & Gia dụng", "doanh_so": 8_500_000_000, "chi_phi_mkt": 680_000_000, "loi_nhuan_rong": 1_190_000_000, "ty_suat_loi_nhuan": 0.140},
    {"nganh_hang": "Thời trang & May mặc", "doanh_so": 4_200_000_000, "chi_phi_mkt": 550_000_000, "loi_nhuan_rong": 840_000_000, "ty_suat_loi_nhuan": 0.200},
    {"nganh_hang": "Mỹ phẩm & Làm đẹp", "doanh_so": 3_800_000_000, "chi_phi_mkt": 600_000_000, "loi_nhuan_rong": 950_000_000, "ty_suat_loi_nhuan": 0.250},
    {"nganh_hang": "Mẹ & Bé", "doanh_so": 2_600_000_000, "chi_phi_mkt": 260_000_000, "loi_nhuan_rong": 416_000_000, "ty_suat_loi_nhuan": 0.160},
    {"nganh_hang": "Bách hóa tiêu dùng", "doanh_so": 5_400_000_000, "chi_phi_mkt": 320_000_000, "loi_nhuan_rong": 486_000_000, "ty_suat_loi_nhuan": 0.090},
]


# --------------------------------------------------------------------------
# Test Cases
# --------------------------------------------------------------------------

def test_pptx_deck_grounding_and_native_shapes(tmp_path):
    """Test 1: Presentation (.pptx) contains native PowerPoint charts and strict grounded numbers."""
    pptx_file = tmp_path / "test_presentation.pptx"
    deck.render_deck(PPTX_SPEC, PPTX_ROWS, pptx_file)
    assert pptx_file.exists() and pptx_file.stat().st_size > 0

    prs = Presentation(pptx_file)
    assert len(prs.slides) == 5

    # Check KPI cards
    kpi_texts = " ".join(shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "12,85 tỷ" in kpi_texts
    assert "27.543" in kpi_texts
    assert "+36,9%" in kpi_texts

    # Check Native Chart Shape
    chart_shapes = [s for s in prs.slides[2].shapes if s.has_chart]
    assert len(chart_shapes) == 1
    chart = chart_shapes[0].chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert [c.label for c in chart.plots[0].categories] == ["Q1", "Q2", "Q3", "Q4"]

    # Check Hallucination Exclusion in Bullets
    bullet_texts = " ".join(shape.text for shape in prs.slides[4].shapes if shape.has_text_frame)
    assert "Kênh Online giữ vai trò động lực" in bullet_texts
    assert "99999999999" not in bullet_texts


def test_xlsx_multi_sheet_growth_and_totals(tmp_path):
    """Test 2: Spreadsheet (.xlsx) contains correct totals sheet and growth formulas."""
    xlsx_file = tmp_path / "test_spreadsheet.xlsx"
    officecli._write_workbook(XLSX_ROWS, xlsx_file, period_column="quy")
    assert xlsx_file.exists()

    wb = load_workbook(xlsx_file)
    assert "data" in wb.sheetnames and "totals" in wb.sheetnames

    totals_dict = dict(wb["totals"].iter_rows(min_row=2, values_only=True))
    assert totals_dict["row_count"] == 12
    assert totals_dict["sum_doanh_thu_thuan"] == 34_550_000_000
    assert totals_dict["sum_chi_phi_van_hanh"] == 15_770_000_000
    assert totals_dict["sum_loi_nhuan_gop"] == 18_780_000_000
    assert totals_dict["sum_so_luong_khach"] == 168_600

    # Test growth computation correctness
    growth_dt = deck.compute_metric(XLSX_ROWS, "growth:doanh_thu_thuan:quy")
    assert totals_dict["growth_doanh_thu_thuan"] == pytest.approx(growth_dt)
    assert totals_dict["growth_doanh_thu_thuan"] == pytest.approx((10.8e9 - 8.7e9) / 8.7e9)


def test_docx_executive_report_traceability(tmp_path):
    """Test 3: Executive Document (.docx) has 100% traceable facts and drops hallucinated numbers."""
    facts = _grounded_facts(DOCX_ROWS)
    assert facts["row_count"] == 5
    assert facts["sum_doanh_so"] == 24_500_000_000
    assert facts["sum_loi_nhuan_rong"] == 3_882_000_000

    prose = """# Báo Cáo Q4 2025
Ngành Điện tử tiếp tục giữ tỷ trọng lớn.
Tổng doanh số bịa đặt đạt 999999999 tỷ đồng.
Ngành Mỹ phẩm đạt tỷ suất sinh lời ấn tượng."""

    cleaned_prose = _strip_unverified_lines(prose)
    assert "999999999" not in cleaned_prose
    assert "Ngành Điện tử tiếp tục giữ tỷ trọng lớn" in cleaned_prose

    full_md = cleaned_prose + "\n\n" + _facts_section(facts)
    html_content = _markdown.markdown(full_md, extensions=["tables"])
    docx_bytes = html_to_docx(html_content, title="Báo Cáo Q4 2025")

    docx_file = tmp_path / "test_report.docx"
    docx_file.write_bytes(docx_bytes)
    assert docx_file.exists()

    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "999999999" not in full_text
    expected_formatted = {deck.format_number(v) for v in facts.values()}

    token_re = re.compile(r"\d[\d.,]*\s*%?")
    for token in token_re.findall(full_text):
        digits = re.sub(r"[.,%\s]", "", token)
        if len(digits) <= 1 or (len(digits) == 4 and 1900 <= int(digits) <= 2099):
            continue
        assert any(token.strip() in exp or exp in token.strip() for exp in expected_formatted)
