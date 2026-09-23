"""Checks for the grounded deck renderer.

All offline — the point of this renderer is that no model touches the numbers,
so nothing here needs one either.
"""

import asyncio
import json
from types import SimpleNamespace

import pytest
from pptx import Presentation

from dbgpt_analyst.tools import deck

ROWS = [
    {"quy": "Q1", "kenh": "Online", "doanh_thu": 1_250_000_000, "so_don": 3120},
    {"quy": "Q1", "kenh": "Offline", "doanh_thu": 890_000_000, "so_don": 1450},
    {"quy": "Q2", "kenh": "Online", "doanh_thu": 1_480_000_000, "so_don": 3600},
    {"quy": "Q2", "kenh": "Offline", "doanh_thu": 910_000_000, "so_don": 1502},
    {"quy": "Q3", "kenh": "Online", "doanh_thu": 1_620_000_000, "so_don": 3980},
    {"quy": "Q3", "kenh": "Offline", "doanh_thu": 870_000_000, "so_don": 1390},
    {"quy": "Q4", "kenh": "Online", "doanh_thu": 2_140_000_000, "so_don": 5210},
    {"quy": "Q4", "kenh": "Offline", "doanh_thu": 1_050_000_000, "so_don": 1710},
]

SPEC = {
    "title": "Doanh thu 2024",
    "subtitle": "Theo quý và kênh",
    "source": "Nguồn: bảng orders",
    "slides": [
        {
            "kind": "kpi",
            "title": "Chỉ số chính",
            "kpis": [
                {"label": "Tổng doanh thu", "metric": "sum:doanh_thu"},
                {"label": "Tổng đơn", "metric": "sum:so_don"},
                {"label": "Tăng trưởng Q4/Q3", "metric": "growth:doanh_thu:quy"},
            ],
        },
        {
            "kind": "chart",
            "title": "Doanh thu theo quý",
            "chart": "column",
            "category_column": "quy",
            "value_column": "doanh_thu",
            "series_column": "kenh",
        },
        {"kind": "table", "title": "Chi tiết", "columns": ["quy", "kenh", "doanh_thu"]},
        {"kind": "bullets", "title": "Nhận định", "bullets": ["Online dẫn dắt tăng trưởng"]},
    ],
}


# --- metrics: the numbers officecli got wrong -----------------------------


def test_growth_is_computed_not_guessed():
    # Q3 totals 2.49bn, Q4 totals 3.19bn. officecli's model reported a growth
    # figure unrelated to its own data; this one is arithmetic.
    assert deck.compute_metric(ROWS, "growth:doanh_thu:quy") == pytest.approx(
        (3_190_000_000 - 2_490_000_000) / 2_490_000_000
    )


def test_basic_aggregates():
    assert deck.compute_metric(ROWS, "sum:doanh_thu") == 10_210_000_000
    assert deck.compute_metric(ROWS, "sum:so_don") == 21_962
    assert deck.compute_metric(ROWS, "count") == 8
    assert deck.compute_metric(ROWS, "max:doanh_thu") == 2_140_000_000
    assert deck.compute_metric(ROWS, "avg:so_don") == pytest.approx(21_962 / 8)


def test_growth_needs_the_period_column_named():
    with pytest.raises(ValueError, match="growth:doanh_thu:quy"):
        deck.compute_metric(ROWS, "growth:doanh_thu")


def test_unknown_column_names_itself(monkeypatch):
    # The spec can come from a model, so a bad column is input, not a crash.
    with pytest.raises(ValueError, match="doanh_thu_sai"):
        deck.compute_metric(ROWS, "sum:doanh_thu_sai")


def test_growth_by_zero_baseline_is_refused():
    rows = [{"p": "A", "v": 0}, {"p": "B", "v": 5}]
    with pytest.raises(ValueError, match="undefined"):
        deck.compute_metric(rows, "growth:v:p")


# --- formatting: the truncation officecli produced ------------------------


def test_vietnamese_number_format_keeps_every_digit():
    # officecli rendered 6.920 as "6." and 28.4% as "28.". Nothing may be lost.
    assert deck.format_number(2_140_000_000) == "2,14 tỷ"
    assert deck.format_number(1_000_000_000) == "1 tỷ"
    assert deck.format_number(1_050_000_000) == "1,05 tỷ"
    assert deck.format_number(6_920) == "6.920"
    assert deck.format_number(3120) == "3.120"


def test_growth_is_formatted_as_a_signed_percent():
    assert deck.format_metric("growth:doanh_thu:quy", 0.281124) == "+28,1%"
    assert deck.format_metric("growth:doanh_thu:quy", -0.05) == "-5,0%"


# --- rendering ------------------------------------------------------------


def test_chart_is_a_native_part_carrying_the_real_values(tmp_path):
    # officecli decks had zero chart objects; ours must be real and editable.
    out = deck.render_deck(SPEC, ROWS, tmp_path / "d.pptx")
    charts = [
        shape.chart
        for slide in Presentation(out).slides
        for shape in slide.shapes
        if shape.has_chart
    ]
    assert len(charts) == 1

    plotted = {s.name: list(s.values) for s in charts[0].plots[0].series}
    assert plotted["Online"] == [1_250_000_000, 1_480_000_000, 1_620_000_000, 2_140_000_000]
    assert plotted["Offline"] == [890_000_000, 910_000_000, 870_000_000, 1_050_000_000]
    assert list(charts[0].plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]


def test_deck_shows_only_numbers_that_came_from_the_rows(tmp_path):
    out = deck.render_deck(SPEC, ROWS, tmp_path / "d.pptx")
    text = "\n".join(
        shape.text_frame.text
        for slide in Presentation(out).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "10,21 tỷ" in text  # the true total
    assert "+28,1%" in text  # the true growth
    assert "21.962" in text


def test_pie_collapses_series_into_one_whole(tmp_path):
    spec = {
        "title": "t",
        "slides": [
            {
                "kind": "chart",
                "title": "Tỷ trọng",
                "chart": "pie",
                "category_column": "kenh",
                "value_column": "doanh_thu",
                "series_column": "quy",
            }
        ],
    }
    out = deck.render_deck(spec, ROWS, tmp_path / "p.pptx")
    chart = next(
        s.chart for sl in Presentation(out).slides for s in sl.shapes if s.has_chart
    )
    series = list(chart.plots[0].series)
    assert len(series) == 1
    assert list(series[0].values) == [3_720_000_000, 6_490_000_000]  # Offline, Online


def test_slide_count_matches_the_spec_plus_a_cover(tmp_path):
    out = deck.render_deck(SPEC, ROWS, tmp_path / "d.pptx")
    assert len(Presentation(out).slides) == len(SPEC["slides"]) + 1


def test_no_foreign_boilerplate_is_appended(tmp_path):
    # officecli always closed with an English "Approve one focused validation
    # cycle" slide that had nothing to do with the data.
    out = deck.render_deck(SPEC, ROWS, tmp_path / "d.pptx")
    text = " ".join(
        shape.text_frame.text
        for slide in Presentation(out).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    for leaked in ("Approve one focused", "Quality Gate", "reaches"):
        assert leaked not in text


def test_bad_slide_kind_is_rejected(tmp_path):
    spec = {"title": "t", "slides": [{"kind": "hologram", "title": "x"}]}
    with pytest.raises(ValueError, match="hologram"):
        deck.render_deck(spec, ROWS, tmp_path / "x.pptx")


def test_empty_rows_are_rejected(tmp_path):
    with pytest.raises(ValueError, match="no rows"):
        deck.render_deck(SPEC, [], tmp_path / "x.pptx")


# --- bullets: the last surface a model could smuggle a number through -----


def _text_of(path):
    return "\n".join(
        shape.text_frame.text
        for slide in Presentation(path).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def _bullets_spec(bullets):
    return {"title": "t", "slides": [{"kind": "bullets", "title": "Nhận định", "bullets": bullets}]}


def test_a_bullet_stating_a_figure_is_dropped(tmp_path):
    out = deck.render_deck(
        _bullets_spec(["Doanh thu tăng 45%", "Online dẫn dắt tăng trưởng"]),
        ROWS,
        tmp_path / "b.pptx",
    )
    text = _text_of(out)
    assert "Online dẫn dắt tăng trưởng" in text
    assert "45%" not in text


def test_labels_that_merely_look_numeric_survive(tmp_path):
    kept = ["Q4 bứt phá", "Năm 2024 khởi sắc", "Top 3 kênh chủ lực", "Kênh Online thắng thế."]
    out = deck.render_deck(_bullets_spec(kept), ROWS, tmp_path / "b.pptx")
    text = _text_of(out)
    for bullet in kept:
        assert bullet in text


@pytest.mark.parametrize(
    "bullet",
    [
        "Tăng trưởng 28,1% so với quý trước",  # a computed figure, unverifiable here
        "Đạt 10,21 tỷ đồng",
        "Giảm 5%",  # a percent is a claim even at one digit
        "Có 21.962 đơn hàng",
    ],
)
def test_every_shape_of_written_number_is_caught(bullet, tmp_path):
    spec = _bullets_spec([bullet, "Kênh Online dẫn dắt"])
    assert bullet not in _text_of(deck.render_deck(spec, ROWS, tmp_path / "b.pptx"))


def test_a_slide_of_nothing_but_written_numbers_is_refused(tmp_path):
    with pytest.raises(ValueError, match="unverified"):
        deck.render_deck(_bullets_spec(["Tăng 45%", "Đạt 500 tỷ"]), ROWS, tmp_path / "b.pptx")


# --- validate_spec: the gate between a model's spec and the renderer ------


def test_validate_spec_drops_a_slide_naming_a_column_that_is_not_there():
    spec = {
        "title": "t",
        "slides": [
            {"kind": "kpi", "title": "Tốt", "kpis": [{"label": "Tổng", "metric": "sum:doanh_thu"}]},
            {
                "kind": "chart",
                "title": "Hỏng",
                "chart": "column",
                "category_column": "quy",
                "value_column": "khong_ton_tai",
            },
        ],
    }
    cleaned, problems = deck.validate_spec(spec, ROWS)
    assert [s["title"] for s in cleaned["slides"]] == ["Tốt"]
    assert len(problems) == 1
    assert "khong_ton_tai" in problems[0]


def test_validate_spec_keeps_a_wholly_good_spec():
    cleaned, problems = deck.validate_spec(SPEC, ROWS)
    assert cleaned["slides"] == SPEC["slides"]
    assert problems == []


def test_validate_spec_survives_a_slide_that_is_not_even_an_object():
    cleaned, problems = deck.validate_spec({"title": "t", "slides": ["rác"]}, ROWS)
    assert cleaned["slides"] == []
    assert len(problems) == 1


# --- default_spec: the floor, derived from the data alone -----------------


def test_default_spec_renders_without_any_model(tmp_path):
    spec = deck.default_spec(ROWS, "Doanh thu")
    out = deck.render_deck(spec, ROWS, tmp_path / "d.pptx")
    presentation = Presentation(out)
    assert len(presentation.slides) == len(spec["slides"]) + 1
    assert any(shape.has_chart for slide in presentation.slides for shape in slide.shapes)
    assert "10,21 tỷ" in _text_of(out)  # the real total, computed here


def test_default_spec_degrades_to_a_table_when_nothing_is_numeric(tmp_path):
    rows = [{"ten": "A", "ghi_chu": "x"}, {"ten": "B", "ghi_chu": "y"}]
    spec = deck.default_spec(rows, "Danh sách")
    assert [s["kind"] for s in spec["slides"]] == ["table"]
    deck.render_deck(spec, rows, tmp_path / "t.pptx")  # renders, does not raise


def test_default_spec_rejects_empty_rows():
    with pytest.raises(ValueError, match="no rows"):
        deck.default_spec([], "t")


# --- plan_deck: a model may pick the story, never the numbers -------------


class _StubLLM:
    """Stands in for the planner: replays one canned reply, or raises."""

    def __init__(self, reply):
        self._reply = reply

    async def ainvoke(self, messages):
        assert any("KHÔNG viết bất kỳ con số nào" in m["content"] for m in messages)
        if isinstance(self._reply, Exception):
            raise self._reply
        return SimpleNamespace(content=self._reply)


def _plan(reply, **kwargs):
    return asyncio.run(deck.plan_deck("Doanh thu 2024", ROWS, llm=_StubLLM(reply), **kwargs))


def test_plan_deck_uses_a_sound_plan_verbatim():
    planned = {
        "title": "Bức tranh doanh thu",
        "slides": [
            {"kind": "kpi", "title": "Chỉ số", "kpis": [{"label": "Tổng", "metric": "sum:doanh_thu"}]},
            {
                "kind": "chart",
                "title": "Theo quý",
                "chart": "line",
                "category_column": "quy",
                "value_column": "doanh_thu",
            },
        ],
    }
    spec = _plan(f"```json\n{json.dumps(planned, ensure_ascii=False)}\n```", source="bảng orders")
    assert spec["title"] == "Bức tranh doanh thu"
    assert [s["title"] for s in spec["slides"]] == ["Chỉ số", "Theo quý"]
    assert spec["source"] == "bảng orders"


def test_plan_deck_keeps_the_good_slides_and_drops_the_rest():
    planned = {
        "slides": [
            {"kind": "kpi", "title": "Giữ", "kpis": [{"label": "Tổng", "metric": "sum:doanh_thu"}]},
            {"kind": "kpi", "title": "Bỏ", "kpis": [{"label": "?", "metric": "sum:bia_ra"}]},
            {"kind": "hologram", "title": "Bỏ nữa"},
        ]
    }
    spec = _plan(json.dumps(planned, ensure_ascii=False))
    assert [s["title"] for s in spec["slides"]] == ["Giữ"]
    assert spec["title"] == "Doanh thu 2024"  # falls back to the question


def test_plan_deck_falls_back_when_the_reply_is_not_json():
    spec = _plan("Xin lỗi, tôi không thể tạo slide.")
    assert spec == deck.default_spec(ROWS, "Doanh thu 2024")


def test_plan_deck_falls_back_when_the_model_itself_fails():
    spec = _plan(RuntimeError("upstream 503"))
    assert spec == deck.default_spec(ROWS, "Doanh thu 2024")


def test_plan_deck_falls_back_when_no_planned_slide_survives():
    planned = {"slides": [{"kind": "kpi", "title": "x", "kpis": [{"label": "?", "metric": "sum:bia"}]}]}
    spec = _plan(json.dumps(planned))
    assert [s["kind"] for s in spec["slides"]] == ["kpi", "chart", "table"]


def test_plan_deck_output_always_renders(tmp_path):
    for reply in ("rác", json.dumps({"slides": [{"kind": "table", "title": "Chi tiết"}]})):
        spec = _plan(reply)
        deck.render_deck(spec, ROWS, tmp_path / "p.pptx")


# --- plan_deck: findings become their own grounded kpi slides -------------


def test_plan_deck_appends_one_slide_per_finding():
    findings = [
        SimpleNamespace(claim="Online tăng mạnh Q4", metric_expr="sum:doanh_thu", value=1234.5),
        {"claim": "Số đơn ổn định", "metric_expr": "avg:so_don", "value": 42.0},
    ]
    spec = _plan("rác", findings=findings)  # planner fails -> default_spec + findings
    finding_slides = spec["slides"][-2:]
    assert [s["title"] for s in finding_slides] == ["Online tăng mạnh Q4", "Số đơn ổn định"]
    assert [s["kpis"][0]["value"] for s in finding_slides] == [1234.5, 42.0]


def test_plan_deck_finding_slide_value_survives_render_untouched_by_compute_metric():
    # DoD: moi so tren slide finding = Finding.value, khong phai compute_metric(rows, ...).
    finding = SimpleNamespace(claim="Giả thuyết lệch dữ liệu", metric_expr="sum:doanh_thu", value=999.0)
    assert deck.compute_metric(ROWS, "sum:doanh_thu") != finding.value

    spec = _plan(json.dumps({"slides": []}), findings=[finding])
    kpi = spec["slides"][-1]["kpis"][0]
    assert kpi["value"] == finding.value

    html = deck.render_html(spec, ROWS)
    assert deck.format_metric("sum:doanh_thu", finding.value) in html


def test_plan_deck_finding_count_holds_across_all_planning_paths():
    findings = [{"claim": "x", "metric_expr": "count", "value": 1.0}]
    for reply in (
        RuntimeError("upstream down"),
        "not json",
        json.dumps({"slides": [{"kind": "kpi", "title": "bad", "kpis": [{"label": "?", "metric": "sum:bia"}]}]}),
        json.dumps({"slides": [{"kind": "table", "title": "Chi tiết"}]}),
    ):
        spec = _plan(reply, findings=findings)
        assert spec["slides"][-1]["title"] == "x"
        assert spec["slides"][-1]["kpis"][0]["value"] == 1.0
