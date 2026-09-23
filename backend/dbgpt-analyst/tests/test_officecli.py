"""Checks for the officecli wrapper.

Everything here is offline except `test_live_report`, which actually spends an
LLM call and is therefore opt-in: set OFFICECLI_LIVE_TEST=1 to run it.

`asyncio.run` rather than pytest-asyncio: no new dependency.
"""

import asyncio
import json
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dbgpt_analyst.tools import officecli
from dbgpt_analyst.tools.deck import compute_metric

ROWS = [
    {"quy": "Q1", "kenh": "Online", "doanh_thu": 1_250_000_000, "so_don": 3120},
    {"quy": "Q1", "kenh": "Offline", "doanh_thu": 890_000_000, "so_don": 1450},
    {"quy": "Q4", "kenh": "Online", "doanh_thu": 2_140_000_000, "so_don": 5210},
]


def test_totals_sheet_precomputes_the_sums_the_model_gets_wrong(tmp_path):
    # officecli was measured reporting 10.28bn for a set that sums to 10.21bn.
    # The totals sheet is the fix, so the sums in it must be exact.
    path = officecli._write_workbook(ROWS, tmp_path / "src.xlsx")
    totals = dict(load_workbook(path)["totals"].iter_rows(min_row=2, values_only=True))

    assert totals["row_count"] == 3
    assert totals["sum_doanh_thu"] == 4_280_000_000
    assert totals["sum_so_don"] == 9780
    assert totals["avg_so_don"] == pytest.approx(3260)


def test_totals_sheet_has_no_growth_rows_without_a_period_column(tmp_path):
    path = officecli._write_workbook(ROWS, tmp_path / "src.xlsx")
    metrics = [r[0] for r in load_workbook(path)["totals"].iter_rows(min_row=2, values_only=True)]
    assert not any(m.startswith("growth_") for m in metrics)


def test_totals_sheet_growth_matches_deck_compute_metric(tmp_path):
    # D2: growth must be computed the exact same way slides compute it, so a
    # report and a deck built from the same rows never disagree.
    path = officecli._write_workbook(ROWS, tmp_path / "src.xlsx", period_column="quy")
    totals = dict(load_workbook(path)["totals"].iter_rows(min_row=2, values_only=True))

    assert totals["growth_doanh_thu"] == pytest.approx(compute_metric(ROWS, "growth:doanh_thu:quy"))
    assert totals["growth_so_don"] == pytest.approx(compute_metric(ROWS, "growth:so_don:quy"))
    # so_don: Q1 4570 -> Q4 5210, a real (non-zero) increase — proves this isn't
    # a coincidental 0 from doanh_thu's Q1==Q4 total.
    assert totals["growth_so_don"] > 0


def test_workbook_keeps_the_raw_rows_too(tmp_path):
    path = officecli._write_workbook(ROWS, tmp_path / "src.xlsx")
    data = load_workbook(path)["data"]
    assert [c.value for c in data[1]] == ["quy", "kenh", "doanh_thu", "so_don"]
    assert data.max_row == len(ROWS) + 1


def test_text_columns_are_not_treated_as_numeric(tmp_path):
    path = officecli._write_workbook(ROWS, tmp_path / "src.xlsx")
    metrics = [r[0] for r in load_workbook(path)["totals"].iter_rows(min_row=2, values_only=True)]
    assert "sum_kenh" not in metrics and "sum_quy" not in metrics


def test_argv_never_uses_interactive_mode(tmp_path):
    # `--mode best` prompts on stdin and hangs a server process forever.
    argv = officecli._build_argv(
        "docx", "topic", tmp_path, prompt=None, data_file=None, lang="vi"
    )
    assert "--mode" in argv and argv[argv.index("--mode") + 1] == "fast"
    assert "--json" in argv


def test_reference_scan_is_only_disabled_where_the_flag_exists(tmp_path):
    # officecli errors out with "--no-reference-scan is only supported for pptx".
    kw = dict(prompt=None, data_file=None, lang="vi")
    assert "--no-reference-scan" in officecli._build_argv("pptx", "t", tmp_path, **kw)
    for kind in ("docx", "xlsx", "report"):
        assert "--no-reference-scan" not in officecli._build_argv(kind, "t", tmp_path, **kw)


def test_rows_reach_the_prompt_when_there_is_no_file_channel():
    prompt = officecli._prompt_with_rows("Bao cao quy 4", ROWS)
    assert "Bao cao quy 4" in prompt
    assert "2140000000" in prompt.replace(",", "").replace(" ", "")
    assert "không bịa" in prompt


def test_long_result_sets_are_capped_but_the_true_size_is_stated():
    many = [{"i": i} for i in range(500)]
    prompt = officecli._prompt_with_rows(None, many)
    assert "60/500" in prompt
    assert json.dumps({"i": 60}) not in prompt


def test_environment_forces_external_runtime(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    monkeypatch.delenv("OFFICE_CLI_LLM_API_KEY", raising=False)
    env = officecli._environment()
    # hosted mode burns metered credits; external must never silently regress.
    assert env["OFFICE_CLI_RUNTIME"] == "external"
    assert env["OFFICE_CLI_LLM_API_KEY"] == "sk-test"


def test_missing_key_fails_loudly_instead_of_falling_back_to_hosted(monkeypatch):
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    monkeypatch.delenv("OFFICE_CLI_LLM_API_KEY", raising=False)
    with pytest.raises(RuntimeError, match="No LLM key"):
        officecli._environment()


@pytest.mark.skipif(
    os.getenv("OFFICECLI_LIVE_TEST") != "1", reason="spends an LLM call; set OFFICECLI_LIVE_TEST=1"
)
def test_live_report(tmp_path):
    out = asyncio.run(
        officecli.generate("report", "Doanh thu theo quy", tmp_path, rows=ROWS)
    )
    assert out.exists() and out.stat().st_size > 0
    html = Path(out).read_text(encoding="utf-8")
    assert "2140" in html.replace(",", "").replace(".", "")
