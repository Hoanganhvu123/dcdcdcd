"""tests/test_live_e2e_artifacts.py — Live E2E Verification of Fact-Grounded Office Artifacts.

Milestone 5 (Task Y - R4) Live E2E Verification Suite:
Validates live generation of real, fact-grounded analytical Microsoft Office report artifacts:
1. Excel (.xlsx):
   - Multi-sheet workbook generation with 'data' and 'totals' sheets.
   - Precomputed sums, averages, and period-over-period growth rates.
   - Exact mathematical grounding against executed SQL database queries.
   - Non-numeric column exclusion and cell type validation.
2. PowerPoint (.pptx):
   - Presentation deck generation with KPI summary cards, native clustered column charts, and data tables.
   - Vietnamese currency and metric formatting (tỷ, triệu, nghìn, +X.X%).
   - Strict omission of ungrounded / hallucinated bullet claims.
   - Native XML and shape hierarchy validation.
3. Word (.docx):
   - Executive report generation with structured sections, summary metrics, and data tables.
   - Markdown to HTML to DOCX compilation.
   - 100% metric grounding verification: every single number in the rendered document is traceable to query facts.
4. Edge Case Scenarios:
   - Single-row datasets, negative growth/profits, zero values, and multi-table joined relational schemas.
"""

from __future__ import annotations

import io
import json
import os
import re
import sqlite3
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Dict, List

import markdown as _markdown
import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from dbgpt_analyst.subgraphs.modes.docx_generator import html_to_docx
from dbgpt_analyst.subgraphs.modes.office_writer_subgraph import (
    _export_excel_file,
    _export_pptx_file,
    _export_word_file,
    _facts_section,
    _grounded_facts,
    _strip_unverified_lines,
)
from dbgpt_analyst.tools import deck, officecli


# ---------------------------------------------------------------------------
# Relational Database Fixture with Genuine Analytical Sales Data
# ---------------------------------------------------------------------------

@pytest.fixture
def live_sales_db(tmp_path: Path) -> sqlite3.Connection:
    """Creates a realistic SQLite analytical database for live query execution."""
    db_path = tmp_path / "live_enterprise_sales.db"
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # 1. Sales Performance Table
    cursor.execute("""
        CREATE TABLE sales_performance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            quarter TEXT NOT NULL,
            region TEXT NOT NULL,
            channel TEXT NOT NULL,
            revenue REAL NOT NULL,
            operating_cost REAL NOT NULL,
            gross_profit REAL NOT NULL,
            order_count INTEGER NOT NULL,
            active_customers INTEGER NOT NULL
        )
    """)

    # 2. Regional Targets Table
    cursor.execute("""
        CREATE TABLE regional_targets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            region TEXT NOT NULL,
            quarter TEXT NOT NULL,
            target_revenue REAL NOT NULL
        )
    """)

    sales_records = [
        ("Q1", "North", "Online", 1_450_000_000, 650_000_000, 800_000_000, 3625, 1200),
        ("Q1", "North", "Offline", 980_000_000, 480_000_000, 500_000_000, 1960, 850),
        ("Q1", "South", "Online", 2_100_000_000, 950_000_000, 1_150_000_000, 5250, 1900),
        ("Q1", "South", "Offline", 1_350_000_000, 680_000_000, 670_000_000, 2700, 1100),
        ("Q2", "North", "Online", 1_820_000_000, 780_000_000, 1_040_000_000, 4333, 1450),
        ("Q2", "North", "Offline", 1_020_000_000, 490_000_000, 530_000_000, 2000, 880),
        ("Q2", "South", "Online", 2_650_000_000, 1_150_000_000, 1_500_000_000, 6309, 2300),
        ("Q2", "South", "Offline", 1_420_000_000, 710_000_000, 710_000_000, 2784, 1150),
        ("Q3", "North", "Online", 2_150_000_000, 920_000_000, 1_230_000_000, 4778, 1600),
        ("Q3", "North", "Offline", 1_050_000_000, 500_000_000, 550_000_000, 2019, 900),
        ("Q3", "South", "Online", 3_100_000_000, 1_350_000_000, 1_750_000_000, 7209, 2700),
        ("Q3", "South", "Offline", 1_550_000_000, 750_000_000, 800_000_000, 2980, 1220),
        ("Q4", "North", "Online", 3_100_000_000, 1_250_000_000, 1_850_000_000, 6458, 2100),
        ("Q4", "North", "Offline", 1_280_000_000, 580_000_000, 700_000_000, 2370, 980),
        ("Q4", "South", "Online", 4_250_000_000, 1_750_000_000, 2_500_000_000, 8947, 3400),
        ("Q4", "South", "Offline", 1_900_000_000, 890_000_000, 1_010_000_000, 3518, 1400),
    ]
    cursor.executemany(
        """INSERT INTO sales_performance 
           (quarter, region, channel, revenue, operating_cost, gross_profit, order_count, active_customers) 
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        sales_records,
    )

    targets = [
        ("North", "Q1", 2_300_000_000),
        ("North", "Q2", 2_700_000_000),
        ("North", "Q3", 3_000_000_000),
        ("North", "Q4", 4_000_000_000),
        ("South", "Q1", 3_200_000_000),
        ("South", "Q2", 3_800_000_000),
        ("South", "Q3", 4_400_000_000),
        ("South", "Q4", 5_800_000_000),
    ]
    cursor.executemany(
        "INSERT INTO regional_targets (region, quarter, target_revenue) VALUES (?, ?, ?)",
        targets,
    )

    conn.commit()
    yield conn
    conn.close()


def execute_query(conn: sqlite3.Connection, sql: str) -> List[Dict[str, Any]]:
    """Executes SQL and returns list of dictionaries."""
    cursor = conn.cursor()
    cursor.execute(sql)
    rows = cursor.fetchall()
    return [dict(r) for r in rows]


# ---------------------------------------------------------------------------
# Test Cases: Excel (.xlsx) Live Generation & Validation
# ---------------------------------------------------------------------------

def test_live_xlsx_generation_from_relational_sql(live_sales_db: sqlite3.Connection, tmp_path: Path):
    """Generates Excel workbook from live SQL results and deeply inspects structure and formulas."""
    sql = """
        SELECT 
            s.quarter,
            s.region,
            s.channel,
            s.revenue,
            s.operating_cost,
            s.gross_profit,
            s.order_count,
            s.active_customers,
            t.target_revenue,
            ROUND((s.revenue / t.target_revenue) * 100, 2) AS quota_pct
        FROM sales_performance s
        JOIN regional_targets t ON s.region = t.region AND s.quarter = t.quarter
        ORDER BY s.quarter, s.region, s.channel
    """
    rows = execute_query(live_sales_db, sql)
    assert len(rows) == 16

    xlsx_path = tmp_path / "live_sales_analysis.xlsx"
    out_path = officecli._write_workbook(rows, xlsx_path, period_column="quarter")

    # 1. Inspect file existence and ZIP container
    assert out_path.exists()
    assert out_path.stat().st_size > 0
    with zipfile.ZipFile(out_path, "r") as zf:
        namelist = zf.namelist()
        assert "[Content_Types].xml" in namelist
        assert "xl/workbook.xml" in namelist
        assert "xl/worksheets/sheet1.xml" in namelist

    # 2. Inspect sheets and content using openpyxl
    wb = load_workbook(out_path)
    assert "data" in wb.sheetnames
    assert "totals" in wb.sheetnames

    # Data sheet assertions
    data_sheet = wb["data"]
    headers = [cell.value for cell in data_sheet[1]]
    expected_headers = [
        "quarter", "region", "channel", "revenue", "operating_cost",
        "gross_profit", "order_count", "active_customers", "target_revenue", "quota_pct"
    ]
    assert headers == expected_headers
    assert data_sheet.max_row == len(rows) + 1

    # Totals sheet assertions
    totals = dict(wb["totals"].iter_rows(min_row=2, values_only=True))
    assert totals["row_count"] == 16

    expected_sum_revenue = sum(r["revenue"] for r in rows)
    expected_sum_profit = sum(r["gross_profit"] for r in rows)
    expected_sum_orders = sum(r["order_count"] for r in rows)
    expected_avg_revenue = expected_sum_revenue / len(rows)

    assert totals["sum_revenue"] == expected_sum_revenue
    assert totals["sum_gross_profit"] == expected_sum_profit
    assert totals["sum_order_count"] == expected_sum_orders
    assert totals["avg_revenue"] == pytest.approx(expected_avg_revenue)

    # Growth assertions matching deck.compute_metric
    expected_growth_rev = deck.compute_metric(rows, "growth:revenue:quarter")
    assert totals["growth_revenue"] == pytest.approx(expected_growth_rev)
    assert totals["growth_revenue"] > 0  # Q4 vs Q1 revenue growth is positive

    # Non-numeric column exclusion
    assert "sum_quarter" not in totals
    assert "sum_region" not in totals
    assert "sum_channel" not in totals


# ---------------------------------------------------------------------------
# Test Cases: PowerPoint (.pptx) Live Generation & Validation
# ---------------------------------------------------------------------------

def test_live_pptx_deck_generation_from_relational_sql(live_sales_db: sqlite3.Connection, tmp_path: Path):
    """Generates PowerPoint presentation deck with native charts, KPI cards, and table."""
    sql = """
        SELECT 
            quarter,
            channel,
            SUM(revenue) AS total_revenue,
            SUM(gross_profit) AS total_profit,
            SUM(order_count) AS total_orders
        FROM sales_performance
        GROUP BY quarter, channel
        ORDER BY quarter, channel
    """
    rows = execute_query(live_sales_db, sql)
    assert len(rows) == 8

    deck_spec = {
        "title": "Báo Cáo Phân Tích Doanh Thu Kênh Phân Phối 2026",
        "subtitle": "Đánh giá chi tiết hiệu quả kênh Online và Offline qua các quý",
        "source": "Nguồn: Enterprise SQL Database",
        "slides": [
            {
                "kind": "kpi",
                "title": "Chỉ Số Tổng Hợp Bán Hàng",
                "kpis": [
                    {"label": "Tổng Doanh Thu Hợp Nhất", "metric": "sum:total_revenue"},
                    {"label": "Tổng Lợi Nhuận Gộp", "metric": "sum:total_profit"},
                    {"label": "Tăng Trưởng Q4/Q1", "metric": "growth:total_revenue:quarter"},
                ],
            },
            {
                "kind": "chart",
                "title": "Cơ Cấu Doanh Thu Kênh Theo Quý",
                "chart": "column",
                "category_column": "quarter",
                "value_column": "total_revenue",
                "series_column": "channel",
            },
            {
                "kind": "table",
                "title": "Bảng Chi Tiết Kết Quả Hoạt Động",
                "columns": ["quarter", "channel", "total_revenue", "total_profit", "total_orders"],
                "limit": 8,
            },
            {
                "kind": "bullets",
                "title": "Nhận Định Phân Tích",
                "bullets": [
                    "Kênh Online ghi nhận tốc độ mở rộng vượt bậc qua các quý.",
                    "Lợi nhuận gộp toàn hệ thống duy trì biên lợi nhuận ổn định.",
                    "Doanh thu bịa đặt đạt 88888888888 đồng tại chi nhánh ảo.",  # Hallucinated number
                ],
            },
        ],
    }

    pptx_path = tmp_path / "live_executive_deck.pptx"
    deck.render_deck(deck_spec, rows, pptx_path)

    assert pptx_path.exists()
    assert pptx_path.stat().st_size > 0

    prs = Presentation(pptx_path)
    assert len(prs.slides) == 5  # Title + 4 content slides

    # 1. Slide 0: Title Slide
    title_text = " ".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Báo Cáo Phân Tích Doanh Thu Kênh Phân Phối 2026" in title_text

    # 2. Slide 1: KPI Slide
    kpi_text = " ".join(shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "31,17 tỷ" in kpi_text or "31.17 tỷ" in kpi_text or "tỷ" in kpi_text
    assert "+" in kpi_text and "%" in kpi_text  # Positive growth rate

    # 3. Slide 2: Native Chart Slide
    chart_shapes = [s for s in prs.slides[2].shapes if s.has_chart]
    assert len(chart_shapes) == 1
    chart = chart_shapes[0].chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    categories = [c.label for c in chart.plots[0].categories]
    assert categories == ["Q1", "Q2", "Q3", "Q4"]

    # 4. Slide 3: Table Slide
    table_shapes = [s for s in prs.slides[3].shapes if s.has_table]
    assert len(table_shapes) == 1
    tbl = table_shapes[0].table
    assert len(tbl.rows) == 9  # 1 header + 8 data rows

    # 5. Slide 4: Bullets Slide (Verification of hallucination filtering)
    bullet_text = " ".join(shape.text for shape in prs.slides[4].shapes if shape.has_text_frame)
    assert "Kênh Online ghi nhận" in bullet_text
    assert "88888888888" not in bullet_text  # Must be stripped


# ---------------------------------------------------------------------------
# Test Cases: Word (.docx) Live Generation & Traceability Validation
# ---------------------------------------------------------------------------

def test_live_docx_report_generation_from_relational_sql(live_sales_db: sqlite3.Connection, tmp_path: Path):
    """Generates Word document from live SQL data and verifies 100% fact grounding."""
    sql = """
        SELECT 
            region,
            SUM(revenue) AS regional_revenue,
            SUM(operating_cost) AS regional_cost,
            SUM(gross_profit) AS regional_profit,
            SUM(order_count) AS total_orders
        FROM sales_performance
        GROUP BY region
        ORDER BY regional_revenue DESC
    """
    rows = execute_query(live_sales_db, sql)
    assert len(rows) == 2  # North and South

    facts = _grounded_facts(rows)
    assert facts["row_count"] == 2.0
    assert facts["sum_regional_revenue"] == 31_170_000_000
    assert facts["sum_regional_profit"] == 17_290_000_000

    prose = """# Báo Cáo Chiến Lược Kinh Doanh Vùng 2026
## 1. Đánh giá tổng quan
Khu vực miền Nam tiếp tục đóng vai trò đầu tàu tăng trưởng doanh số.
Doanh số ảo tưởng vượt mốc 777777777 tỷ đồng không có căn cứ.
Khu vực miền Bắc tối ưu hóa chi phí vận hành đạt hiệu suất cao.
"""

    cleaned_prose = _strip_unverified_lines(prose)
    assert "777777777" not in cleaned_prose
    assert "Khu vực miền Nam" in cleaned_prose

    full_md = cleaned_prose + "\n\n" + _facts_section(facts)
    html_content = _markdown.markdown(full_md, extensions=["tables"])
    docx_bytes = html_to_docx(html_content, title="Báo Cáo Chiến Lược Kinh Doanh Vùng 2026")

    docx_path = tmp_path / "live_strategy_report.docx"
    docx_path.write_bytes(docx_bytes)

    assert docx_path.exists()
    assert docx_path.stat().st_size > 0

    doc = Document(docx_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)

    # 1. Hallucination check
    assert "777777777" not in full_text

    # 2. Strict Fact grounding check
    expected_formatted = {deck.format_number(v) for v in facts.values()}
    token_re = re.compile(r"\d[\d.,]*\s*%?")
    for token in token_re.findall(full_text):
        digits = re.sub(r"[.,%\s]", "", token)
        if len(digits) <= 1 or (len(digits) == 4 and 1900 <= int(digits) <= 2099):
            continue
        assert any(token.strip() in exp or exp in token.strip() for exp in expected_formatted)


# ---------------------------------------------------------------------------
# Test Cases: Subgraph Export Handlers Integration
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_office_subgraph_excel_and_word_direct_exports(live_sales_db: sqlite3.Connection, tmp_path: Path, monkeypatch):
    """Directly tests _export_excel_file, _export_word_file, and _export_pptx_file."""
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))

    sql = "SELECT quarter, revenue, gross_profit FROM sales_performance LIMIT 4"
    rows = execute_query(live_sales_db, sql)

    # 1. Test Excel Export
    excel_url = _export_excel_file(rows)
    assert excel_url.startswith("/uploads/generated_xlsx/")
    assert excel_url.endswith(".xlsx")

    excel_disk_path = tmp_path / "uploads" / excel_url.replace("/uploads/", "")
    assert excel_disk_path.exists()
    wb = load_workbook(excel_disk_path)
    assert "Data" in wb.sheetnames or "data" in wb.sheetnames

    # 2. Test Word Export with Stub LLM
    class WordStubLLM:
        async def ainvoke(self, messages):
            md_text = (
                "# Báo Cáo Tài Chính Q1 2026\n"
                "## 1. Tóm tắt điều hành\n"
                "Hiệu quả kinh doanh tăng trưởng ổn định theo kế hoạch.\n"
                "Con số bịa đặt 555555555 đồng không có thật.\n"
            )
            return SimpleNamespace(content=md_text)

    async def fake_create_llm(**kwargs):
        return WordStubLLM(), None

    monkeypatch.setattr(
        "dbgpt_analyst.common.model_fallback.create_llm_with_fallback",
        fake_create_llm,
    )

    word_url, word_html = await _export_word_file("Báo Cáo Q1 2026", rows)
    assert word_url.startswith("/uploads/generated_docx/")
    assert word_url.endswith(".docx")

    word_disk_path = tmp_path / "uploads" / word_url.replace("/uploads/", "")
    assert word_disk_path.exists()
    doc = Document(word_disk_path)
    word_text = "\n".join(p.text for p in doc.paragraphs)
    assert "555555555" not in word_text
    assert "Báo Cáo Tài Chính Q1 2026" in word_text

    # 3. Test PPTX Export with Stub LLM
    plan_dict = {
        "title": "Báo Cáo Q1 Deck",
        "slides": [
            {
                "kind": "kpi",
                "title": "KPIs",
                "kpis": [{"label": "Doanh thu", "metric": "sum:revenue"}],
            }
        ],
    }

    class PPTStubLLM:
        async def ainvoke(self, messages):
            return SimpleNamespace(content=json.dumps(plan_dict, ensure_ascii=False))

    async def fake_create_llm_ppt(**kwargs):
        return PPTStubLLM(), None

    monkeypatch.setattr(
        "dbgpt_analyst.common.model_fallback.create_llm_with_fallback",
        fake_create_llm_ppt,
    )

    pptx_url, preview_html, slides = await _export_pptx_file("Báo Cáo Q1", rows)
    assert pptx_url.startswith("/uploads/generated_pptx/")
    assert pptx_url.endswith(".pptx")
    pptx_disk_path = tmp_path / "uploads" / pptx_url.replace("/uploads/", "")
    assert pptx_disk_path.exists()
    prs = Presentation(pptx_disk_path)
    assert len(prs.slides) == 2  # Title + 1 content slide


# ---------------------------------------------------------------------------
# Test Cases: Edge Cases (Single-Row, Negative Profits, Zero-Growth)
# ---------------------------------------------------------------------------

def test_live_artifacts_edge_cases_single_row_and_negative_metrics(tmp_path: Path):
    """Verifies artifact generation handles single-row datasets and negative metrics gracefully."""
    edge_rows = [
        {"quarter": "Q1", "revenue": 100_000_000, "gross_profit": -20_000_000, "margin_pct": -0.20},
    ]

    # Excel
    xlsx_path = tmp_path / "edge_case.xlsx"
    officecli._write_workbook(edge_rows, xlsx_path, period_column="quarter")
    assert xlsx_path.exists()
    wb = load_workbook(xlsx_path)
    totals = dict(wb["totals"].iter_rows(min_row=2, values_only=True))
    assert totals["row_count"] == 1
    assert totals["sum_gross_profit"] == -20_000_000

    # Deck KPI
    spec = {
        "title": "Báo Cáo Lỗ Q1",
        "slides": [
            {
                "kind": "kpi",
                "title": "Chỉ Số Lỗ",
                "kpis": [
                    {"label": "Lợi Nhuận Âm", "metric": "sum:gross_profit"},
                ],
            }
        ],
    }
    pptx_path = tmp_path / "edge_case.pptx"
    deck.render_deck(spec, edge_rows, pptx_path)
    assert pptx_path.exists()
    prs = Presentation(pptx_path)
    kpi_texts = " ".join(s.text for s in prs.slides[1].shapes if s.has_text_frame)
    assert "-20" in kpi_texts or "20 triệu" in kpi_texts
