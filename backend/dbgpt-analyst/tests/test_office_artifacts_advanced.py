"""Comprehensive verification test suite for Advanced Office Artifacts Generation:
1. Multi-Sheet Spreadsheet (.xlsx) Engine with interconnected formulas & professional styling.
2. Executive Presentation Slide Deck (.pptx) Engine with callout, takeaway, comparison, and native editable charts.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from dbgpt_analyst.subgraphs.modes import office_writer_subgraph as ow
from dbgpt_analyst.tools import deck, officecli

SAMPLE_FINANCIAL_ROWS = [
    {"quy": "Q1", "kenh": "Online", "doanh_thu": 1_250_000_000, "chi_phi": 750_000_000, "so_don": 3120},
    {"quy": "Q1", "kenh": "Offline", "doanh_thu": 890_000_000, "chi_phi": 620_000_000, "so_don": 1450},
    {"quy": "Q2", "kenh": "Online", "doanh_thu": 1_480_000_000, "chi_phi": 810_000_000, "so_don": 3600},
    {"quy": "Q2", "kenh": "Offline", "doanh_thu": 910_000_000, "chi_phi": 600_000_000, "so_don": 1502},
    {"quy": "Q3", "kenh": "Online", "doanh_thu": 1_620_000_000, "chi_phi": 890_000_000, "so_don": 3980},
    {"quy": "Q3", "kenh": "Offline", "doanh_thu": 870_000_000, "chi_phi": 580_000_000, "so_don": 1390},
    {"quy": "Q4", "kenh": "Online", "doanh_thu": 2_140_000_000, "chi_phi": 1_100_000_000, "so_don": 5210},
    {"quy": "Q4", "kenh": "Offline", "doanh_thu": 1_050_000_000, "chi_phi": 690_000_000, "so_don": 1710},
]


# ==============================================================================
# 1. Multi-Sheet Spreadsheet (.xlsx) Engine Tests
# ==============================================================================


def test_export_excel_file_generates_multisheet_workbook_with_formulas(tmp_path: Path, monkeypatch):
    """Verify that _export_excel_file generates a 3-sheet workbook with native formulas."""
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    url = ow._export_excel_file(SAMPLE_FINANCIAL_ROWS)

    assert url is not None
    assert url.startswith("/uploads/generated_xlsx/")
    assert url.endswith(".xlsx")

    file_path = tmp_path / "uploads" / url.replace("/uploads/", "")
    assert file_path.exists()

    wb = load_workbook(file_path, data_only=False)
    # Check all three interconnected sheets exist
    assert "Executive Summary" in wb.sheetnames
    assert "Data" in wb.sheetnames
    assert "KPI Metrics" in wb.sheetnames

    # Check Data sheet row count exactly matches rows + 1 (header)
    assert wb["Data"].max_row == len(SAMPLE_FINANCIAL_ROWS) + 1

    # Check Executive Summary contains native Excel formulas referencing Data
    ws_exec = wb["Executive Summary"]
    formulas_found = []
    for row in ws_exec.iter_rows(values_only=False):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas_found.append(cell.value)

    assert len(formulas_found) >= 5
    # Verify presence of native formula families
    assert any("COUNTA(Data!" in f for f in formulas_found)
    assert any("SUM(Data!" in f for f in formulas_found)
    assert any("AVERAGE(Data!" in f for f in formulas_found)
    assert any("MAX(Data!" in f for f in formulas_found)
    assert any("MIN(Data!" in f for f in formulas_found)
    assert any("IF(Data!" in f for f in formulas_found)

    # Check KPI Metrics sheet formulas
    ws_kpi = wb["KPI Metrics"]
    kpi_formulas = []
    for row in ws_kpi.iter_rows(values_only=False):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                kpi_formulas.append(cell.value)

    assert len(kpi_formulas) >= 3
    assert any("SUMIF(Data!" in f for f in kpi_formulas)
    assert any("AVERAGEIF(Data!" in f for f in kpi_formulas)
    assert any("SUM(" in f for f in kpi_formulas)


def test_officecli_create_multi_sheet_workbook(tmp_path: Path):
    """Verify create_multi_sheet_workbook in officecli."""
    out_path = tmp_path / "multisheet_officecli.xlsx"
    officecli.create_multi_sheet_workbook(SAMPLE_FINANCIAL_ROWS, out_path, period_column="quy")

    assert out_path.exists()
    wb = load_workbook(out_path, data_only=False)
    assert "Executive Summary" in wb.sheetnames
    assert "data" in wb.sheetnames
    assert "totals" in wb.sheetnames
    assert "KPI Metrics" in wb.sheetnames

    # Verify formulas in Executive Summary
    exec_formulas = [
        c.value for row in wb["Executive Summary"].iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert any("SUM(data!" in f for f in exec_formulas)
    assert any("AVERAGE(data!" in f for f in exec_formulas)


def test_render_excel_html_preview_multi_sheet_tabs():
    """Verify HTML spreadsheet preview produces interactive tabs and formula bindings."""
    html = ow._render_excel_html("Báo cáo doanh thu và chi phí 2024", SAMPLE_FINANCIAL_ROWS)
    assert "<!DOCTYPE html>" in html
    assert "Executive Summary" in html
    assert "KPI Metrics" in html
    assert "sheet-exec" in html
    assert "sheet-data" in html
    assert "sheet-kpi" in html
    assert 'data-formula="=COUNTA(Data!A2:A9)"' in html
    assert "#1E3A8A" in html  # Dark navy header color styling


# ==============================================================================
# 2. Executive Presentation Slide Deck (.pptx) Engine Tests
# ==============================================================================


def test_deck_callout_component_rendering(tmp_path: Path):
    """Verify callout slide renders container, accent highlight bar, headline, and insights."""
    spec = {
        "title": "Báo cáo Điều hành Chiến lược",
        "slides": [
            {
                "kind": "callout",
                "title": "Tâm điểm Thị trường",
                "subtitle": "Phân tích động lực tăng trưởng",
                "badge": "ĐỘT PHÁ DOANH THU",
                "headline": "Kênh Online bứt phá dẫn dắt toàn hệ thống",
                "insights": [
                    "Quy mô doanh thu kênh Online tăng trưởng liên tục qua 4 quý.",
                    "Hiệu quả chi phí vận hành được kiểm soát tối ưu.",
                ],
            }
        ],
    }
    out = deck.render_deck(spec, SAMPLE_FINANCIAL_ROWS, tmp_path / "callout.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 2  # Cover + Callout slide

    slide = prs.slides[1]
    # Check shapes: title textbox, subtitle textbox, rule, container rounded rect, accent bar, badge, headline, insights
    assert len(slide.shapes) >= 6

    text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Kênh Online bứt phá dẫn dắt toàn hệ thống" in text
    assert "ĐỘT PHÁ DOANH THU" in text
    assert "Quy mô doanh thu kênh Online" in text

    # HTML preview verification
    html = deck.render_html(spec, SAMPLE_FINANCIAL_ROWS)
    assert "callout-box" in html
    assert "callout-headline" in html
    assert "Kênh Online bứt phá" in html


def test_deck_takeaway_component_rendering(tmp_path: Path):
    """Verify takeaway slide renders structured priority action plan cards."""
    spec = {
        "title": "Kế hoạch Hành động",
        "slides": [
            {
                "kind": "takeaway",
                "title": "Đề xuất & Khuyến nghị Chiến lược",
                "subtitle": "Lộ trình triển khai 3 trọng tâm",
                "takeaways": [
                    {
                        "priority": "Ưu tiên 1",
                        "title": "Mở rộng kênh bán hàng trực tuyến",
                        "description": "Tối ưu ngân sách quảng cáo và cá nhân hóa trải nghiệm người dùng.",
                    },
                    {
                        "priority": "Ưu tiên 2",
                        "title": "Cải tiến quy trình vận hành chuỗi cung ứng",
                        "description": "Rút ngắn thời gian giao hàng và giảm tỷ lệ hoàn hủy.",
                    },
                    {
                        "priority": "Ưu tiên 3",
                        "title": "Tự động hóa quản trị tài chính",
                        "description": "Ứng dụng AI phân tích dữ liệu theo thời gian thực.",
                    },
                ],
            }
        ],
    }
    out = deck.render_deck(spec, SAMPLE_FINANCIAL_ROWS, tmp_path / "takeaway.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 2

    slide = prs.slides[1]
    text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Mở rộng kênh bán hàng trực tuyến" in text
    assert "Cải tiến quy trình vận hành" in text
    assert "Tự động hóa quản trị tài chính" in text
    assert "Ưu tiên 1" in text

    html = deck.render_html(spec, SAMPLE_FINANCIAL_ROWS)
    assert "takeaway-grid" in html
    assert "takeaway-card" in html
    assert "Mở rộng kênh bán hàng" in html


def test_deck_comparison_component_rendering(tmp_path: Path):
    """Verify comparison slide renders multi-column comparison cards with calculated metrics."""
    spec = {
        "title": "Đánh giá Phân khúc",
        "slides": [
            {
                "kind": "comparison",
                "title": "So sánh Hiệu quả Kênh Phân phối",
                "subtitle": "Kênh Online vs Kênh Offline",
                "columns": [
                    {
                        "header": "Kênh Trực Tuyến (Online)",
                        "badge": "Tăng trưởng nhanh",
                        "metric": "sum:doanh_thu",
                        "bullets": [
                            "Đóng góp tỷ trọng lớn vào doanh số",
                            "Tốc độ bứt phá mạnh mẽ",
                        ],
                    },
                    {
                        "header": "Kênh Truyền Thống (Offline)",
                        "badge": "Ổn định",
                        "bullets": [
                            "Mạng lưới đại lý phủ rộng",
                            "Chi phí mặt bằng cố định cần tối ưu",
                        ],
                    },
                ],
            }
        ],
    }
    out = deck.render_deck(spec, SAMPLE_FINANCIAL_ROWS, tmp_path / "comparison.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 2

    slide = prs.slides[1]
    text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Kênh Trực Tuyến (Online)" in text
    assert "Kênh Truyền Thống (Offline)" in text
    assert "Tăng trưởng nhanh" in text
    assert "10,21 tỷ" in text  # Verified computed sum:doanh_thu

    html = deck.render_html(spec, SAMPLE_FINANCIAL_ROWS)
    assert "comp-grid" in html
    assert "comp-card" in html
    assert "10,21 tỷ" in html


def test_native_editable_chart_shapes_all_four_types(tmp_path: Path):
    """Verify all 4 chart types (Column, Bar, Line, Pie) render native PowerPoint charts."""
    spec = {
        "title": "Bộ biểu đồ đa dạng",
        "slides": [
            {
                "kind": "chart",
                "title": "Column Clustered Chart",
                "chart": "column",
                "category_column": "quy",
                "value_column": "doanh_thu",
                "series_column": "kenh",
            },
            {
                "kind": "chart",
                "title": "Bar Clustered Chart",
                "chart": "bar",
                "category_column": "quy",
                "value_column": "chi_phi",
            },
            {
                "kind": "chart",
                "title": "Line Markers Chart",
                "chart": "line",
                "category_column": "quy",
                "value_column": "so_don",
            },
            {
                "kind": "chart",
                "title": "Pie Distribution Chart",
                "chart": "pie",
                "category_column": "kenh",
                "value_column": "doanh_thu",
            },
        ],
    }
    out = deck.render_deck(spec, SAMPLE_FINANCIAL_ROWS, tmp_path / "all_charts.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 5  # Cover + 4 chart slides

    expected_types = [
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        XL_CHART_TYPE.BAR_CLUSTERED,
        XL_CHART_TYPE.LINE_MARKERS,
        XL_CHART_TYPE.PIE,
    ]

    for i, exp_type in enumerate(expected_types, start=1):
        slide = prs.slides[i]
        charts = [s.chart for s in slide.shapes if s.has_chart]
        assert len(charts) == 1
        assert charts[0].chart_type == exp_type
        # Verify raw data points are populated in series
        plot = charts[0].plots[0]
        assert len(plot.series) >= 1
        assert len(list(plot.categories)) >= 2


def test_comprehensive_executive_narrative_deck(tmp_path: Path):
    """Verify build_executive_deck_spec produces a complete 7-8 slide executive deck."""
    spec = deck.build_executive_deck_spec(
        SAMPLE_FINANCIAL_ROWS,
        "Báo cáo Điều hành Kết quả Kinh doanh Toàn diện",
        subtitle="Phân tích Doanh thu & Chi phí 2024",
        source="Cơ sở dữ liệu Bán hàng Doanh nghiệp",
    )
    slides = spec["slides"]
    assert 6 <= len(slides) <= 10

    slide_kinds = [s["kind"] for s in slides]
    assert "kpi" in slide_kinds
    assert "chart" in slide_kinds
    assert "callout" in slide_kinds
    assert "table" in slide_kinds
    assert "comparison" in slide_kinds
    assert "takeaway" in slide_kinds

    # Render to pptx file and in-memory bytes
    out_path = deck.render_deck(spec, SAMPLE_FINANCIAL_ROWS, tmp_path / "executive_deck.pptx")
    assert out_path.exists()
    prs = Presentation(out_path)
    assert len(prs.slides) == len(slides) + 1

    deck_bytes = deck.render_deck_bytes(spec, SAMPLE_FINANCIAL_ROWS)
    assert len(deck_bytes) > 10_000
    prs_from_bytes = Presentation(io.BytesIO(deck_bytes))
    assert len(prs_from_bytes.slides) == len(slides) + 1

    # Validate HTML preview
    html = deck.render_html(spec, SAMPLE_FINANCIAL_ROWS)
    assert "Báo cáo Điều hành" in html
    assert "callout-box" in html
    assert "takeaway-grid" in html
    assert "comp-grid" in html
