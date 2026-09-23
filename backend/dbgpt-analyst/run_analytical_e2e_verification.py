"""Comprehensive E2E Analytical Grounding & Office Artifact Verification.

Executes 3 realistic Vietnamese business analytical scenarios:
1. Presentation (.pptx): Omnichannel sales performance with native PPT charts and grounded KPI cards.
2. Spreadsheet (.xlsx): Multi-sheet financial/operational workbook with precomputed growth and totals.
3. Executive Document (.docx): Strategic business unit report with strict factual grounding.

Verifies:
- Mathematical accuracy of metrics (sums, averages, growth %).
- Preservation of native chart objects & editable shapes in PPTX.
- Accurate totals & growth sheet formulas in XLSX.
- Strict rejection of unverified model hallucinations in DOCX.
- Saves all files to backend/dbgpt-analyst/test_outputs/.
"""

import io
import json
import os
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

# Import dbgpt_analyst modules
from dbgpt_analyst.tools import deck, officecli
from dbgpt_analyst.subgraphs.modes.docx_generator import html_to_docx
from dbgpt_analyst.subgraphs.modes.office_writer_subgraph import (
    _grounded_facts,
    _strip_unverified_lines,
    _facts_section,
    _export_excel_file,
)
import markdown as _markdown

OUTPUT_DIR = Path(__file__).resolve().parent / "test_outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# --------------------------------------------------------------------------
# Scenario 1: Omnichannel Retail Performance (.pptx)
# --------------------------------------------------------------------------
PPTX_ROWS = [
    {"quy": "Q1", "kenh": "Online", "doanh_thu": 1_450_000_000, "so_don": 3625, "aov": 400_000},
    {"quy": "Q1", "kenh": "Offline", "doanh_thu": 980_000_000, "so_don": 1960, "aov": 500_000},
    {"quy": "Q2", "kenh": "Online", "doanh_thu": 1_820_000_000, "so_don": 4333, "aov": 420_000},
    {"quy": "Q2", "kenh": "Offline", "doanh_thu": 1_020_000_000, "so_don": 2000, "aov": 510_000},
    {"quy": "Q3", "kenh": "Online", "doanh_thu": 2_150_000_000, "so_don": 4778, "aov": 450_000},
    {"quy": "Q3", "kenh": "Offline", "doanh_thu": 1_050_000_000, "so_don": 2019, "aov": 520_000},
    {"quy": "Q4", "kenh": "Online", "doanh_thu": 3_100_000_000, "so_don": 6458, "aov": 480_000},
    {"quy": "Q4", "kenh": "Offline", "doanh_thu": 1_280_000_000, "so_don": 2370, "aov": 540_000},
]

PPTX_SPEC = {
    "title": "Báo Cáo Phân Tích Doanh Thu Kênh Phân Phối 2025",
    "subtitle": "Đánh giá tăng trưởng các quý và hiệu quả kênh Online so với Offline",
    "source": "Nguồn: Hệ thống ERP & CRM Doanh Nghiệp Bán Lẻ",
    "slides": [
        {
            "kind": "kpi",
            "title": "Chỉ Số Hiệu Suất Tổng Quan",
            "kpis": [
                {"label": "Tổng doanh thu toàn quốc", "metric": "sum:doanh_thu"},
                {"label": "Tổng số lượng đơn hàng", "metric": "sum:so_don"},
                {"label": "Tăng trưởng doanh thu Q4/Q3", "metric": "growth:doanh_thu:quy"},
            ],
        },
        {
            "kind": "chart",
            "title": "Cơ Cấu Doanh Thu Theo Quý & Kênh",
            "chart": "column",
            "category_column": "quy",
            "value_column": "doanh_thu",
            "series_column": "kenh",
        },
        {
            "kind": "table",
            "title": "Bảng Dữ Liệu Bán Hàng Chi Tiết",
            "columns": ["quy", "kenh", "doanh_thu", "so_don", "aov"],
            "limit": 8,
        },
        {
            "kind": "bullets",
            "title": "Nhận Định Chiến Lược & Xu Hướng",
            "bullets": [
                "Kênh Online giữ vai trò động lực tăng trưởng chính qua các quý",
                "Kênh Offline duy trì doanh số ổn định và giá trị đơn hàng trung bình cao",
                "Quý 4 ghi nhận mức bứt phá doanh số mạnh mẽ nhờ các chiến dịch mùa lễ hội",
                "Doanh thu đạt mức kỷ lục 99999999999 đồng vào cuối năm",  # Fabricated number, must be stripped!
            ],
        },
    ],
}


def verify_presentation_pptx():
    print("\n" + "=" * 70)
    print("RUN 1: Presentation (.pptx) Omnichannel Performance Verification")
    print("=" * 70)

    # 1. Grounding calculations check
    sum_dt = deck.compute_metric(PPTX_ROWS, "sum:doanh_thu")
    sum_don = deck.compute_metric(PPTX_ROWS, "sum:so_don")
    growth_q4_q3 = deck.compute_metric(PPTX_ROWS, "growth:doanh_thu:quy")

    expected_sum_dt = 12_850_000_000.0
    expected_sum_don = 27_543.0
    q3_total = 2_150_000_000 + 1_050_000_000
    q4_total = 3_100_000_000 + 1_280_000_000
    expected_growth = (q4_total - q3_total) / q3_total

    assert sum_dt == expected_sum_dt, f"Expected {expected_sum_dt}, got {sum_dt}"
    assert sum_don == expected_sum_don, f"Expected {expected_sum_don}, got {sum_don}"
    assert abs(growth_q4_q3 - expected_growth) < 1e-6, f"Expected {expected_growth}, got {growth_q4_q3}"
    print(f"[*] Mathematical Grounding OK: Sum Doanh Thu = {deck.format_number(sum_dt)}, Sum Đơn = {deck.format_number(sum_don)}, Growth Q4/Q3 = {deck.format_metric('growth:doanh_thu:quy', growth_q4_q3)}")

    # 2. Render Presentation
    pptx_path = OUTPUT_DIR / "presentation_omnichannel_2025.pptx"
    deck.render_deck(PPTX_SPEC, PPTX_ROWS, pptx_path)
    assert pptx_path.exists() and pptx_path.stat().st_size > 0
    print(f"[*] Rendered PPTX to: {pptx_path} (Size: {pptx_path.stat().st_size} bytes)")

    # 3. HTML Preview Generation
    html_preview = deck.render_html(PPTX_SPEC, PPTX_ROWS)
    html_preview_path = OUTPUT_DIR / "presentation_omnichannel_2025_preview.html"
    html_preview_path.write_text(html_preview, encoding="utf-8")
    assert "12,85 tỷ" in html_preview
    assert "27.543" in html_preview
    assert "+36,9%" in html_preview
    assert "99999999999" not in html_preview  # Stripped!
    print(f"[*] Rendered HTML Preview to: {html_preview_path} (Length: {len(html_preview)} chars)")

    # 4. Deep Inspection of PPTX Structure via python-pptx
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 5, f"Expected 5 slides (Cover + 4 content slides), got {len(prs.slides)}"

    # Slide 1: Cover
    cover_texts = [shape.text for shape in prs.slides[0].shapes if shape.has_text_frame]
    assert any("Báo Cáo Phân Tích Doanh Thu Kênh Phân Phối 2025" in t for t in cover_texts)

    # Slide 2: KPI Cards
    kpi_texts = " ".join(shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "12,85 tỷ" in kpi_texts
    assert "27.543" in kpi_texts
    assert "+36,9%" in kpi_texts
    print("[*] Slide 2 (KPIs): Verified grounded cards for Revenue, Orders, and Growth %")

    # Slide 3: Chart
    chart_shapes = [s for s in prs.slides[2].shapes if s.has_chart]
    assert len(chart_shapes) == 1, "Slide 3 must contain a native PowerPoint chart shape"
    chart = chart_shapes[0].chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    categories = [c.label for c in chart.plots[0].categories]
    assert categories == ["Q1", "Q2", "Q3", "Q4"], f"Chart categories mismatch: {categories}"
    series_names = [s.name for s in chart.series]
    assert series_names == ["Offline", "Online"] or series_names == ["Online", "Offline"]
    print("[*] Slide 3 (Native Chart): Verified editable Column Clustered chart with 4 quarterly categories and 2 channel series")

    # Slide 4: Table
    table_shapes = [s for s in prs.slides[3].shapes if s.has_table]
    assert len(table_shapes) == 1, "Slide 4 must contain a native PowerPoint table shape"
    table = table_shapes[0].table
    assert len(table.columns) == 5
    assert len(table.rows) == len(PPTX_ROWS) + 1  # header + 8 rows
    print("[*] Slide 4 (Native Table): Verified 5 columns x 9 rows with exact ERP records")

    # Slide 5: Bullets (Grounded Qualitative Prose)
    bullet_texts = " ".join(shape.text for shape in prs.slides[4].shapes if shape.has_text_frame)
    assert "Kênh Online giữ vai trò động lực tăng trưởng" in bullet_texts
    assert "Quý 4 ghi nhận mức bứt phá" in bullet_texts
    assert "99999999999" not in bullet_texts, "Unverified hallucinated number was NOT stripped!"
    print("[*] Slide 5 (Bullets): Verified hallucinated numbers were strictly stripped")

    print("[SUCCESS] Run 1: Presentation (.pptx) passed all verification gates.\n")


# --------------------------------------------------------------------------
# Scenario 2: Multi-Region Financial & Operational Workbook (.xlsx)
# --------------------------------------------------------------------------
XLSX_ROWS = [
    {"vung_mien": "Miền Bắc", "quy": "Q1", "doanh_thu_thuan": 2_500_000_000, "chi_phi_van_hanh": 1_200_000_000, "loi_nhuan_gop": 1_300_000_000, "so_luong_khach": 12_500},
    {"vung_mien": "Miền Trung", "quy": "Q1", "doanh_thu_thuan": 1_400_000_000, "chi_phi_van_hanh": 750_000_000, "loi_nhuan_gop": 650_000_000, "so_luong_khach": 7_200},
    {"vung_mien": "Miền Nam", "quy": "Q1", "doanh_thu_thuan": 3_200_000_000, "chi_phi_van_hanh": 1_500_000_000, "loi_nhuan_gop": 1_700_000_000, "so_luong_khach": 16_000},
    {"vung_mien": "Miền Bắc", "quy": "Q2", "doanh_thu_thuan": 2_800_000_000, "chi_phi_van_hanh": 1_300_000_000, "loi_nhuan_gop": 1_500_000_000, "so_luong_khach": 13_800},
    {"vung_mien": "Miền Trung", "quy": "Q2", "doanh_thu_thuan": 1_550_000_000, "chi_phi_van_hanh": 800_000_000, "loi_nhuan_gop": 750_000_000, "so_luong_khach": 7_800},
    {"vung_mien": "Miền Nam", "quy": "Q2", "doanh_thu_thuan": 3_600_000_000, "chi_phi_van_hanh": 1_650_000_000, "loi_nhuan_gop": 1_950_000_000, "so_luong_khach": 17_500},
    {"vung_mien": "Miền Bắc", "quy": "Q3", "doanh_thu_thuan": 3_100_000_000, "chi_phi_van_hanh": 1_400_000_000, "loi_nhuan_gop": 1_700_000_000, "so_luong_khach": 15_000},
    {"vung_mien": "Miền Trung", "quy": "Q3", "doanh_thu_thuan": 1_650_000_000, "chi_phi_van_hanh": 820_000_000, "loi_nhuan_gop": 830_000_000, "so_luong_khach": 8_100},
    {"vung_mien": "Miền Nam", "quy": "Q3", "doanh_thu_thuan": 3_950_000_000, "chi_phi_van_hanh": 1_750_000_000, "loi_nhuan_gop": 2_200_000_000, "so_luong_khach": 19_000},
    {"vung_mien": "Miền Bắc", "quy": "Q4", "doanh_thu_thuan": 3_900_000_000, "chi_phi_van_hanh": 1_650_000_000, "loi_nhuan_gop": 2_250_000_000, "so_luong_khach": 18_500},
    {"vung_mien": "Miền Trung", "quy": "Q4", "doanh_thu_thuan": 2_100_000_000, "chi_phi_van_hanh": 950_000_000, "loi_nhuan_gop": 1_150_000_000, "so_luong_khach": 10_200},
    {"vung_mien": "Miền Nam", "quy": "Q4", "doanh_thu_thuan": 4_800_000_000, "chi_phi_van_hanh": 2_000_000_000, "loi_nhuan_gop": 2_800_000_000, "so_luong_khach": 23_000},
]


def verify_spreadsheet_xlsx():
    print("\n" + "=" * 70)
    print("RUN 2: Multi-Sheet Spreadsheet (.xlsx) Verification")
    print("=" * 70)

    xlsx_path = OUTPUT_DIR / "spreadsheet_financial_operational_2025.xlsx"
    officecli._write_workbook(XLSX_ROWS, xlsx_path, period_column="quy")
    assert xlsx_path.exists() and xlsx_path.stat().st_size > 0
    print(f"[*] Generated XLSX workbook: {xlsx_path} (Size: {xlsx_path.stat().st_size} bytes)")

    # Inspect with openpyxl
    wb = load_workbook(xlsx_path)
    assert set(wb.sheetnames) == {"data", "totals"}, f"Expected sheets ['data', 'totals'], got {wb.sheetnames}"

    # 1. Verify 'data' sheet
    ws_data = wb["data"]
    headers = [cell.value for cell in ws_data[1]]
    expected_headers = ["vung_mien", "quy", "doanh_thu_thuan", "chi_phi_van_hanh", "loi_nhuan_gop", "so_luong_khach"]
    assert headers == expected_headers, f"Headers mismatch: {headers}"
    assert ws_data.max_row == len(XLSX_ROWS) + 1  # 13 rows total
    print(f"[*] Sheet 'data': Verified 6 headers and {len(XLSX_ROWS)} row records")

    # 2. Verify 'totals' sheet
    ws_totals = wb["totals"]
    totals_dict = dict(ws_totals.iter_rows(min_row=2, values_only=True))

    expected_sum_dt = sum(r["doanh_thu_thuan"] for r in XLSX_ROWS)
    expected_sum_cp = sum(r["chi_phi_van_hanh"] for r in XLSX_ROWS)
    expected_sum_ln = sum(r["loi_nhuan_gop"] for r in XLSX_ROWS)
    expected_sum_kh = sum(r["so_luong_khach"] for r in XLSX_ROWS)

    assert totals_dict["row_count"] == 12
    assert totals_dict["sum_doanh_thu_thuan"] == expected_sum_dt == 34_550_000_000
    assert totals_dict["sum_chi_phi_van_hanh"] == expected_sum_cp == 15_770_000_000
    assert totals_dict["sum_loi_nhuan_gop"] == expected_sum_ln == 18_780_000_000
    assert totals_dict["sum_so_luong_khach"] == expected_sum_kh == 168_600

    assert totals_dict["avg_doanh_thu_thuan"] == pytest.approx(expected_sum_dt / 12)
    assert totals_dict["avg_chi_phi_van_hanh"] == pytest.approx(expected_sum_cp / 12)
    assert totals_dict["avg_loi_nhuan_gop"] == pytest.approx(expected_sum_ln / 12)
    assert totals_dict["avg_so_luong_khach"] == pytest.approx(expected_sum_kh / 12)

    # 3. Growth metrics check (Q4 vs Q3)
    growth_dt = deck.compute_metric(XLSX_ROWS, "growth:doanh_thu_thuan:quy")
    growth_cp = deck.compute_metric(XLSX_ROWS, "growth:chi_phi_van_hanh:quy")
    growth_ln = deck.compute_metric(XLSX_ROWS, "growth:loi_nhuan_gop:quy")
    growth_kh = deck.compute_metric(XLSX_ROWS, "growth:so_luong_khach:quy")

    assert totals_dict["growth_doanh_thu_thuan"] == pytest.approx(growth_dt)
    assert totals_dict["growth_chi_phi_van_hanh"] == pytest.approx(growth_cp)
    assert totals_dict["growth_loi_nhuan_gop"] == pytest.approx(growth_ln)
    assert totals_dict["growth_so_luong_khach"] == pytest.approx(growth_kh)

    # Validate math: Q4 revenue = 3.9B+2.1B+4.8B = 10.8B, Q3 revenue = 3.1B+1.65B+3.95B = 8.7B
    assert growth_dt == pytest.approx((10.8e9 - 8.7e9) / 8.7e9)
    print(f"[*] Sheet 'totals': Verified precomputed formulas: Sum DT={totals_dict['sum_doanh_thu_thuan']:,} VND, Growth DT={totals_dict['growth_doanh_thu_thuan']:.2%}")

    # Also test xlsx export with chart via _export_excel_file
    excel_url = _export_excel_file(XLSX_ROWS)
    assert excel_url and excel_url.endswith(".xlsx")
    print(f"[*] Exported workbook via _export_excel_file to: {excel_url}")

    print("[SUCCESS] Run 2: Spreadsheet (.xlsx) passed all verification gates.\n")


# --------------------------------------------------------------------------
# Scenario 3: Executive Document (.docx) Grounding Verification
# --------------------------------------------------------------------------
DOCX_ROWS = [
    {"nganh_hang": "Điện tử & Gia dụng", "doanh_so": 8_500_000_000, "chi_phi_mkt": 680_000_000, "loi_nhuan_rong": 1_190_000_000, "ty_suat_loi_nhuan": 0.140},
    {"nganh_hang": "Thời trang & May mặc", "doanh_so": 4_200_000_000, "chi_phi_mkt": 550_000_000, "loi_nhuan_rong": 840_000_000, "ty_suat_loi_nhuan": 0.200},
    {"nganh_hang": "Mỹ phẩm & Làm đẹp", "doanh_so": 3_800_000_000, "chi_phi_mkt": 600_000_000, "loi_nhuan_rong": 950_000_000, "ty_suat_loi_nhuan": 0.250},
    {"nganh_hang": "Mẹ & Bé", "doanh_so": 2_600_000_000, "chi_phi_mkt": 260_000_000, "loi_nhuan_rong": 416_000_000, "ty_suat_loi_nhuan": 0.160},
    {"nganh_hang": "Bách hóa tiêu dùng", "doanh_so": 5_400_000_000, "chi_phi_mkt": 320_000_000, "loi_nhuan_rong": 486_000_000, "ty_suat_loi_nhuan": 0.090},
]


def verify_executive_document_docx():
    print("\n" + "=" * 70)
    print("RUN 3: Executive Document (.docx) Grounding Verification")
    print("=" * 70)

    # 1. Compute grounded facts
    facts = _grounded_facts(DOCX_ROWS)
    assert facts["row_count"] == 5
    assert facts["sum_doanh_so"] == 24_500_000_000.0
    assert facts["sum_chi_phi_mkt"] == 2_410_000_000.0
    assert facts["sum_loi_nhuan_rong"] == 3_882_000_000.0
    assert facts["avg_ty_suat_loi_nhuan"] == pytest.approx(0.168)
    print(f"[*] Computed Grounded Facts: Sum Doanh Số = {deck.format_number(facts['sum_doanh_so'])}, Sum Lợi Nhuận = {deck.format_number(facts['sum_loi_nhuan_rong'])}")

    # 2. Simulate model prose output containing both valid qualitative insights and hallucinated raw numbers
    simulated_model_output = """# Báo Cáo Chiến Lược Hiệu Quả Kinh Doanh Q4 2025

## 1. Tóm Tắt Điều Hành
Ngành hàng Điện tử & Gia dụng tiếp tục dẫn đầu về quy mô doanh số trong danh mục sản phẩm.
Ngành Mỹ phẩm & Làm đẹp đạt tỷ suất sinh lời vượt trội nhờ tối ưu hóa chi phí quảng cáo số.
Tổng doanh số toàn quốc đạt 999999999 tỷ đồng, tăng trưởng 8888% theo ước tính không căn cứ.

## 2. Nhận Định Cơ Cấu Chi Phí & Lợi Nhuận
Chi phí marketing được kiểm soát chặt chẽ ở nhóm Bách hóa tiêu dùng.
Thời trang & May mặc duy trì biên lợi nhuận ổn định qua các kênh phân phối chính thức.
Dự báo quý tới lợi nhuận tăng thêm 777777777 đồng từ thị trường ngách."""

    # 3. Strip unverified lines (the grounding filter)
    cleaned_prose = _strip_unverified_lines(simulated_model_output)
    facts_appendix = _facts_section(facts)
    full_markdown = (cleaned_prose + "\n\n" + facts_appendix).strip()

    assert "999999999" not in cleaned_prose, "Hallucinated number 999999999 must be stripped"
    assert "8888" not in cleaned_prose, "Hallucinated percentage 8888 must be stripped"
    assert "777777777" not in cleaned_prose, "Hallucinated number 777777777 must be stripped"
    assert "Ngành hàng Điện tử & Gia dụng tiếp tục dẫn đầu" in cleaned_prose
    assert "Ngành Mỹ phẩm & Làm đẹp đạt tỷ suất sinh lời" in cleaned_prose
    print("[*] Verified prose filter: All 3 hallucinated lines successfully removed; qualitative insights preserved")

    # 4. Generate docx binary
    html_content = _markdown.markdown(full_markdown, extensions=["tables"])
    docx_bytes = html_to_docx(html_content, title="Báo Cáo Chiến Lược Hiệu Quả Kinh Doanh Q4 2025")

    docx_path = OUTPUT_DIR / "executive_report_strategic_q4_2025.docx"
    docx_path.write_bytes(docx_bytes)
    assert docx_path.exists() and docx_path.stat().st_size > 0
    print(f"[*] Generated DOCX report to: {docx_path} (Size: {docx_path.stat().st_size} bytes)")

    # 5. Parse and audit every token in the generated DOCX file
    doc = Document(io.BytesIO(docx_bytes))
    full_doc_text = "\n".join(p.text for p in doc.paragraphs)

    expected_formatted_numbers = {deck.format_number(v) for v in facts.values()}
    number_token_pattern = re.compile(r"\d[\d.,]*\s*%?")

    for token in number_token_pattern.findall(full_doc_text):
        digits = re.sub(r"[.,%\s]", "", token)
        if len(digits) <= 1:
            continue  # Single digit like section '1', '2', or 'Q4'
        if len(digits) == 4 and 1900 <= int(digits) <= 2099:
            continue  # Year '2025'
        assert any(token.strip() in exp or exp in token.strip() for exp in expected_formatted_numbers), (
            f"Unverified number token '{token}' in DOCX is not in grounded facts: {expected_formatted_numbers}"
        )

    print("[*] Strict Traceability Audit: 100% of numbers in DOCX verified against DB query facts")
    print("[SUCCESS] Run 3: Executive Document (.docx) passed all verification gates.\n")


def main():
    print("=" * 70)
    print("STARTING COMPREHENSIVE E2E ANALYTICAL TEST RUNS & GROUNDING VERIFICATION")
    print("=" * 70)
    verify_presentation_pptx()
    verify_spreadsheet_xlsx()
    verify_executive_document_docx()
    print("=" * 70)
    print("ALL 3 ANALYTICAL RUNS SUCCESSFULLY EXECUTED AND VERIFIED 100% GROUNDED!")
    print("=" * 70)


if __name__ == "__main__":
    main()
