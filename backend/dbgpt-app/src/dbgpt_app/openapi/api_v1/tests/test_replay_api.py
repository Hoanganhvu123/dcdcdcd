"""src/dbgpt_app/openapi/api_v1/tests/test_replay_api.py

Co-located test suite for DB-GPT Replay Subsystem & Presentation Deck Generation.
"""
import base64
import io
import json
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from dbgpt_app.openapi.api_v1.analyst_api import router as analyst_router
from dbgpt_app.openapi.api_v1.deck_service import generate_deck_presentation
from dbgpt_app.openapi.api_v1.replay_api import router as replay_router
from dbgpt_app.openapi.api_v1.replay_schemas import (
    DeckGenerateRequest,
    DeckGenerateResponse,
    ReplaySessionDetailResponse,
    ReplaySessionListResponse,
    SlideKpiItem,
    SlideSpec,
)
from dbgpt_app.openapi.api_v1.seed_replay_data import get_replay_repository


@pytest.fixture(scope="module")
def app() -> FastAPI:
    """Create test FastAPI application with mounted routers."""
    test_app = FastAPI(title="DB-GPT Replay Test App")
    test_app.include_router(replay_router)
    test_app.include_router(analyst_router)
    return test_app


@pytest.fixture(scope="module")
def client(app: FastAPI) -> TestClient:
    """FastAPI TestClient fixture."""
    return TestClient(app)


# ═══════════════════════════════════════════════════════════════════════════
# 1. Replay Sessions Index & Filtering Tests
# ═══════════════════════════════════════════════════════════════════════════

def test_list_sessions_default(client: TestClient):
    """Test listing replay sessions with default pagination."""
    response = client.get("/api/v1/replay/sessions")
    assert response.status_code == 200
    data = response.json()

    assert data["code"] == 0
    assert data["message"] == "success"
    assert "data" in data
    assert data["data"]["total"] >= 3
    assert data["data"]["page"] == 1
    assert len(data["data"]["items"]) >= 3

    # Validate first summary record structure
    first = data["data"]["items"][0]
    assert "sessionId" in first
    assert "title" in first
    assert "userQuery" in first
    assert "status" in first
    assert "totalTurns" in first
    assert "totalSteps" in first
    assert "artifactCount" in first
    assert "previewTags" in first


def test_list_sessions_pagination(client: TestClient):
    """Test page size and total pages calculation."""
    res_page_1 = client.get("/api/v1/replay/sessions?page=1&pageSize=2")
    assert res_page_1.status_code == 200
    d1 = res_page_1.json()["data"]
    assert d1["page"] == 1
    assert d1["pageSize"] == 2
    assert len(d1["items"]) == 2
    assert d1["totalPages"] >= 2

    res_page_2 = client.get("/api/v1/replay/sessions?page=2&pageSize=2")
    assert res_page_2.status_code == 200
    d2 = res_page_2.json()["data"]
    assert d2["page"] == 2
    assert len(d2["items"]) >= 1

    # Ensure items on page 1 and page 2 are distinct
    p1_ids = {item["sessionId"] for item in d1["items"]}
    p2_ids = {item["sessionId"] for item in d2["items"]}
    assert p1_ids.isdisjoint(p2_ids)


def test_list_sessions_out_of_bounds(client: TestClient):
    """Test requesting a page beyond available items returns empty list."""
    res = client.get("/api/v1/replay/sessions?page=999&pageSize=20")
    assert res.status_code == 200
    d = res.json()["data"]
    assert d["page"] == 999
    assert d["items"] == []
    assert d["total"] >= 3


def test_list_sessions_search_filter(client: TestClient):
    """Test case-insensitive search filter."""
    # Search for Canifa
    res = client.get("/api/v1/replay/sessions?search=Canifa")
    assert res.status_code == 200
    items = res.json()["data"]["items"]
    assert len(items) >= 1
    assert any("canifa" in item["sessionId"].lower() or "canifa" in item["title"].lower() for item in items)

    # Search for non-existent keyword
    res_none = client.get("/api/v1/replay/sessions?search=NonExistentKeywordXYZ123")
    assert res_none.status_code == 200
    assert res_none.json()["data"]["total"] == 0
    assert len(res_none.json()["data"]["items"]) == 0


def test_list_sessions_artifact_filter(client: TestClient):
    """Test filtering sessions by artifact kind."""
    res_slide = client.get("/api/v1/replay/sessions?artifactKind=slide")
    assert res_slide.status_code == 200
    items = res_slide.json()["data"]["items"]
    assert len(items) >= 1
    # Check that canifa session is returned
    sids = [item["sessionId"] for item in items]
    assert "canifa-sales-q3-deepdive" in sids


# ═══════════════════════════════════════════════════════════════════════════
# 2. Replay Session Detail & 404 Error Tests
# ═══════════════════════════════════════════════════════════════════════════

def test_get_session_detail_canifa(client: TestClient):
    """Test retrieving full detail of Canifa Sales Q3 Deep-Dive session."""
    res = client.get("/api/v1/replay/session/canifa-sales-q3-deepdive")
    assert res.status_code == 200
    data = res.json()
    assert data["code"] == 0
    sess = data["data"]
    assert sess["sessionId"] == "canifa-sales-q3-deepdive"
    assert "Canifa" in sess["title"]
    assert sess["mode"] == "slides"
    assert sess["status"] == "completed"

    # Verify multi-turn structure (4 turns)
    assert len(sess["turns"]) == 4
    turn1 = sess["turns"][0]
    assert turn1["turnIndex"] == 1
    assert "doanh thu q3" in turn1["query"].lower()
    assert len(turn1["steps"]) >= 2
    assert turn1["sqlUsed"] is not None

    # Verify thinking block in turn 1
    step1 = turn1["steps"][0]
    assert step1["thinking"] is not None
    assert step1["thinking"]["tokens"] > 0
    assert len(step1["thinking"]["subThoughtItems"]) >= 2

    # Verify tool calls
    step2 = turn1["steps"][1]
    assert len(step2["toolCalls"]) >= 1
    assert step2["toolCalls"][0]["toolType"] == "sql"

    # Verify Turn 4 deck artifact
    turn4 = sess["turns"][3]
    assert len(turn4["artifacts"]) >= 1
    art = turn4["artifacts"][0]
    assert art["artifactKind"] == "slide"
    assert art["slideCount"] == 6


def test_get_session_detail_inventory(client: TestClient):
    """Test retrieving Inventory Stockout Forecast session detail."""
    res = client.get("/api/v1/replay/session/inventory-stockout-forecast")
    assert res.status_code == 200
    sess = res.json()["data"]
    assert sess["sessionId"] == "inventory-stockout-forecast"
    assert len(sess["turns"]) == 3
    assert sess["mode"] == "sheets"


def test_get_session_detail_customer_churn(client: TestClient):
    """Test retrieving Customer Churn Cohort session detail."""
    res = client.get("/api/v1/replay/session/customer-churn-cohort")
    assert res.status_code == 200
    sess = res.json()["data"]
    assert sess["sessionId"] == "customer-churn-cohort"
    assert len(sess["turns"]) == 3
    assert sess["mode"] == "deep-research"


def test_get_session_detail_not_found(client: TestClient):
    """Test 404 response for unknown session ID."""
    res = client.get("/api/v1/replay/session/non-existent-session-uuid-999")
    assert res.status_code == 404
    detail = res.json()["detail"]
    assert "not found" in detail.lower()


# ═══════════════════════════════════════════════════════════════════════════
# 3. Replay SSE Wire Streaming Tests
# ═══════════════════════════════════════════════════════════════════════════

def test_stream_replay_session_events(client: TestClient):
    """Test SSE streaming endpoint with instant speed multiplier."""
    response = client.get("/api/v1/replay/session/canifa-sales-q3-deepdive/stream?speed=0")
    assert response.status_code == 200
    assert "text/event-stream" in response.headers["content-type"]
    assert response.headers["cache-control"] == "no-cache"

    raw_text = response.text
    assert len(raw_text) > 0

    # Parse SSE events
    events = []
    current_event = None
    for line in raw_text.splitlines():
        if line.startswith("event:"):
            current_event = line.replace("event:", "").strip()
        elif line.startswith("data:") and current_event:
            data_str = line.replace("data:", "").strip()
            try:
                parsed_data = json.loads(data_str)
                events.append((current_event, parsed_data))
            except Exception:
                pass
            current_event = None

    assert len(events) >= 10

    event_types = [ev[0] for ev in events]
    assert "session_start" in event_types
    assert "turn_start" in event_types
    assert "step_start" in event_types
    assert "thought_chunk" in event_types
    assert "tool_call" in event_types
    assert "tool_result" in event_types
    assert "artifact_created" in event_types
    assert "turn_end" in event_types
    assert "session_end" in event_types


def test_stream_replay_session_not_found(client: TestClient):
    """Test streaming unknown session returns 404."""
    res = client.get("/api/v1/replay/session/unknown-session-id/stream")
    assert res.status_code == 404


# ═══════════════════════════════════════════════════════════════════════════
# 4. Presentation Deck Generation Tests
# ═══════════════════════════════════════════════════════════════════════════

def test_deck_generate_endpoint_analyst(client: TestClient):
    """Test POST /api/v1/analyst/deck/generate with structured slides & query data."""
    req_body = {
        "title": "Báo Cáo Hiệu Quả Kinh Doanh Q3 Canifa",
        "subtitle": "Phân tích doanh thu và tối ưu hóa chiết khấu chuỗi bán lẻ",
        "theme": "emerald_minimal",
        "data": [
            {"quy": "Q1", "doanh_thu": 18400000, "loi_nhuan": 3680000},
            {"quy": "Q2", "doanh_thu": 21020000, "loi_nhuan": 4204000},
            {"quy": "Q3", "doanh_thu": 24850000, "loi_nhuan": 4970000},
        ],
        "slides": [
            {
                "slideNumber": 1,
                "layoutType": "title",
                "title": "Báo Cáo Hiệu Quả Kinh Doanh Q3 Canifa",
                "subtitle": "Phân tích doanh thu và tối ưu hóa chiết khấu chuỗi bán lẻ"
            },
            {
                "slideNumber": 2,
                "layoutType": "metric_cards",
                "title": "Chỉ Số Tài Chính Q3",
                "kpis": [
                    {"label": "Tổng doanh thu", "metric": "sum:doanh_thu"},
                    {"label": "Tăng trưởng Q3/Q2", "metric": "growth:doanh_thu:quy"}
                ]
            },
            {
                "slideNumber": 3,
                "layoutType": "comparison",
                "title": "So Sánh Hiệu Quả Vùng",
                "columns": [
                    {"header": "Hà Nội", "badge": "DẪN ĐẦU", "bullets": ["Doanh số ổn định", "Chiết khấu kiểm soát tốt"]},
                    {"header": "TP.HCM", "badge": "TĂNG TRƯỞNG", "bullets": ["Doanh số online tăng trưởng", "Mở rộng thêm showroom mới"]}
                ]
            },
            {
                "slideNumber": 4,
                "layoutType": "callout",
                "title": "Cảnh Báo Kiểm Toán",
                "headline": "Outlet Hà Đông có tỷ lệ chiết khấu vượt trần quy định",
                "insights": ["Nguyên nhân: giảm giá kép", "Hành động: khóa giảm giá kép trên POS"]
            },
            {
                "slideNumber": 5,
                "layoutType": "takeaway",
                "title": "Kế Hoạch Hành Động Q4",
                "takeaways": [
                    {"priority": "01", "title": "Chuẩn hóa POS", "description": "Thiết lập trần chiết khấu 25%."},
                    {"priority": "02", "title": "Đẩy mạnh E-commerce", "description": "Tập trung chuyển đổi số."}
                ]
            }
        ]
    }

    response = client.post("/api/v1/analyst/deck/generate", json=req_body)
    assert response.status_code == 200
    data = response.json()

    assert data["status"] == "success"
    assert "presentationId" in data
    assert data["title"] == "Báo Cáo Hiệu Quả Kinh Doanh Q3 Canifa"
    assert data["slideCount"] >= 5
    assert data["base64Content"] is not None
    assert len(data["base64Content"]) > 1000
    assert "<!doctype html>" in data["htmlPreview"].lower() or "<div" in data["htmlPreview"] or "<html" in data["htmlPreview"]

    # Validate that base64 decoded bytes form a real, valid PPTX presentation
    pptx_bytes = base64.b64decode(data["base64Content"])
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) >= 5


def test_deck_generate_endpoint_replay_alias(client: TestClient):
    """Test POST /api/v1/replay/deck/generate alias endpoint with auto-derived deck."""
    req_body = {
        "title": "Báo Cáo Phân Tích Chuỗi Cung Ứng",
        "subtitle": "Dự báo tồn kho và tối ưu hóa điều chuyển",
        "theme": "executive_navy",
        "data": [
            {"kho": "Kho Bắc", "ton_kho": 45000, "so_ngay": 120},
            {"kho": "Kho Nam", "ton_kho": 12000, "so_ngay": 14},
        ]
    }

    response = client.post("/api/v1/replay/deck/generate", json=req_body)
    assert response.status_code == 200
    data = response.json()

    assert data["status"] == "success"
    assert data["presentationId"] is not None
    assert data["slideCount"] >= 3
    assert data["base64Content"] is not None


def test_deck_download_endpoint(client: TestClient):
    """Test downloading generated .pptx file and 404 for missing file."""
    # First generate a deck
    req = DeckGenerateRequest(title="Downloadable Deck Test")
    deck_res = generate_deck_presentation(req)
    filename = f"{deck_res.presentation_id}.pptx"

    # Download it
    dl_res = client.get(f"/api/v1/analyst/deck/download/{filename}")
    assert dl_res.status_code == 200
    assert dl_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert len(dl_res.content) > 1000

    # Verify opening the downloaded bytes
    prs = Presentation(io.BytesIO(dl_res.content))
    assert len(prs.slides) >= 3

    # Test downloading non-existent file
    dl_res_404 = client.get("/api/v1/analyst/deck/download/non_existent_file_999.pptx")
    assert dl_res_404.status_code == 404
